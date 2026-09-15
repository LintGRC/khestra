"""Statement of Applicability store — per-control status for ISO 27001.

Mirrors the SoA philosophy of https://github.com/kriss-b/llm-iso27001:
every control starts as "not implemented" and is deliberately reviewed.
"""

from __future__ import annotations

import csv
import io
import json
import os
import sqlite3
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional

from controls import ALL_CONTROLS
from attributes_data import CONTROL_ATTRIBUTES

STATUSES = ("implemented", "partially implemented", "not implemented", "excluded")

_db_path: str = ""


def _db() -> sqlite3.Connection:
    db = sqlite3.connect(_db_path)
    db.row_factory = sqlite3.Row
    return db


def init_store(data_dir: str) -> None:
    global _db_path
    _db_path = os.path.join(data_dir, "soa.db")
    db = _db()
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS soa (
          control_id TEXT PRIMARY KEY,
          section TEXT NOT NULL DEFAULT '',
          title TEXT NOT NULL DEFAULT '',
          summary TEXT NOT NULL DEFAULT '',
          status TEXT NOT NULL DEFAULT 'not implemented',
          applicable INTEGER NOT NULL DEFAULT 1,
          justification TEXT NOT NULL DEFAULT '',
          doc_link TEXT NOT NULL DEFAULT '',
          updated_at TEXT NOT NULL DEFAULT ''
        )
        """
    )
    cols = {r[1] for r in db.execute("PRAGMA table_info(soa)").fetchall()}
    if "attributes" not in cols:
        db.execute("ALTER TABLE soa ADD COLUMN attributes TEXT NOT NULL DEFAULT '{}'")
    db.commit()
    for c in ALL_CONTROLS:
        db.execute(
            "INSERT OR IGNORE INTO soa (control_id, section, title, summary) VALUES (?, ?, ?, ?)",
            (c["id"], c["section"], c["title"], c["summary"]),
        )
    # backfill official 27002 attributes for rows that don't have them yet
    for cid, attrs in CONTROL_ATTRIBUTES.items():
        db.execute(
            "UPDATE soa SET attributes = ? WHERE control_id = ? AND attributes = '{}'",
            (json.dumps(attrs), cid),
        )
    db.commit()
    db.close()


def _row_to_dict(row: sqlite3.Row) -> Dict[str, Any]:
    d = dict(row)
    d["applicable"] = bool(d["applicable"])
    attrs = d.get("attributes")
    if isinstance(attrs, str):
        try:
            d["attributes"] = json.loads(attrs) if attrs else {}
        except (json.JSONDecodeError, TypeError):
            d["attributes"] = {}
    return d


def list_soa() -> List[Dict[str, Any]]:
    db = _db()
    rows = db.execute(
        "SELECT * FROM soa ORDER BY section, control_id"
    ).fetchall()
    db.close()
    return [_row_to_dict(r) for r in rows]


def rollup() -> Dict[str, Any]:
    db = _db()
    rows = db.execute("SELECT section, status, applicable FROM soa").fetchall()
    db.close()
    counts: Dict[str, int] = {s: 0 for s in STATUSES}
    by_section: Dict[str, Dict[str, int]] = {}
    applicable_total = 0
    for r in rows:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
        sec = r["section"]
        if sec not in by_section:
            by_section[sec] = {"total": 0, "implemented": 0}
        by_section[sec]["total"] += 1
        if r["status"] == "implemented":
            by_section[sec]["implemented"] += 1
        if r["applicable"]:
            applicable_total += 1
    return {
        "counts": counts,
        "total": len(rows),
        "applicable_total": applicable_total,
        "excluded": counts.get("excluded", 0),
        "by_section": by_section,
    }


def get_control(control_id: str) -> Optional[Dict[str, Any]]:
    db = _db()
    row = db.execute("SELECT * FROM soa WHERE control_id = ?", (control_id,)).fetchone()
    db.close()
    return _row_to_dict(row) if row else None


def update_control(control_id: str, fields: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    current = get_control(control_id)
    if not current:
        return None
    sets = []
    params: List[Any] = []

    new_status = fields["status"] if "status" in fields else current["status"]
    if "status" in fields and new_status not in STATUSES:
        raise ValueError(f"Invalid status: {new_status}")

    if "applicable" in fields:
        new_applicable = bool(fields["applicable"])
    elif new_status == "excluded":
        new_applicable = False
    elif current["status"] == "excluded" and new_status != "excluded":
        new_applicable = True
    else:
        new_applicable = bool(current["applicable"])
    if new_status == "excluded":
        new_applicable = False

    new_justification = (
        str(fields["justification"]) if "justification" in fields
        else str(current.get("justification") or "")
    )
    if (new_status == "excluded" or not new_applicable) and not new_justification.strip():
        raise ValueError(
            "Exclusion requires a justification (ISO/IEC 27001:2022 clause 6.1.3)."
        )

    if new_status != current["status"]:
        sets.append("status = ?")
        params.append(new_status)
    if int(new_applicable) != int(bool(current["applicable"])):
        sets.append("applicable = ?")
        params.append(1 if new_applicable else 0)
    if "justification" in fields:
        sets.append("justification = ?")
        params.append(new_justification)
    if "doc_link" in fields:
        sets.append("doc_link = ?")
        params.append(str(fields["doc_link"] or ""))
    if "attributes" in fields:
        sets.append("attributes = ?")
        params.append(json.dumps(fields["attributes"] or {}))
    if not sets:
        return current

    sets.append("updated_at = ?")
    params.append(datetime.now(timezone.utc).isoformat(timespec="seconds"))
    params.append(control_id)
    db = _db()
    db.execute(f"UPDATE soa SET {', '.join(sets)} WHERE control_id = ?", params)
    db.commit()
    db.close()
    return get_control(control_id)


def export_csv() -> str:
    db = _db()
    rows = db.execute(
        "SELECT control_id, section, title, status, applicable, justification, doc_link, attributes FROM soa ORDER BY section, control_id"
    ).fetchall()
    db.close()
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(
        ["control_id", "section", "title", "status", "applicable", "justification", "doc_link", "attributes"]
    )
    for r in rows:
        attrs = r["attributes"] or "{}"
        try:
            parsed = json.loads(attrs)
        except (json.JSONDecodeError, TypeError):
            parsed = {}
        writer.writerow([r["control_id"], r["section"], r["title"], r["status"],
                         r["applicable"], r["justification"], r["doc_link"],
                         " | ".join(f"{k}: {', '.join(v)}" for k, v in parsed.items() if v)])
    return "\ufeff" + buf.getvalue()


def _rows_for_export() -> List[Dict[str, Any]]:
    db = _db()
    rows = db.execute(
        "SELECT control_id, section, title, summary, status, applicable, justification, doc_link, attributes "
        "FROM soa ORDER BY section, control_id"
    ).fetchall()
    db.close()
    out = []
    for r in rows:
        d = dict(r)
        d["applicable"] = bool(d["applicable"])
        attrs = d.get("attributes") or "{}"
        try:
            parsed = json.loads(attrs) if isinstance(attrs, str) else {}
        except (json.JSONDecodeError, TypeError):
            parsed = {}
        d["attributes_summary"] = " | ".join(
            f"{k}: {', '.join(v)}" for k, v in parsed.items() if v
        )
        d["attributes"] = parsed
        out.append(d)
    return out


def export_xlsx() -> bytes:
    """ISO-style Statement of Applicability workbook (93 Annex A controls)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    rows = _rows_for_export()
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement of Applicability"

    headers = ["Control", "Section", "Title", "Summary", "Status", "Applicable",
               "Justification", "Doc Link", "ISO 27002 Attributes"]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for r in rows:
        ws.append([
            r["control_id"],
            r["section"],
            r["title"],
            r["summary"],
            r["status"],
            "Yes" if r["applicable"] else "No",
            r["justification"],
            r["doc_link"],
            r["attributes_summary"],
        ])

    widths = [12, 12, 42, 46, 20, 11, 40, 28, 48]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_docx() -> bytes:
    """Management-ready Statement of Applicability document (93 Annex A controls)."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt

    rows = _rows_for_export()
    summary = rollup()

    doc = Document()
    doc.add_heading("ISO/IEC 27001:2022 — Statement of Applicability", 0)
    doc.add_paragraph(
        "This Statement of Applicability records the implementation status of every "
        "Annex A control for the in-scope information security management system."
    )

    counts = summary["counts"]
    doc.add_heading("Summary", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Metric"
    hdr[1].text = "Value"
    for label, value in [
        ("Total controls (Annex A 2022)", summary["total"]),
        ("Applicable", summary["applicable_total"]),
        ("Implemented", counts.get("implemented", 0)),
        ("Partially implemented", counts.get("partially implemented", 0)),
        ("Not implemented", counts.get("not implemented", 0)),
        ("Excluded", counts.get("excluded", 0)),
    ]:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = str(value)

    doc.add_heading("Controls", level=1)
    soa_table = doc.add_table(rows=1, cols=6)
    soa_table.style = "Table Grid"
    soa_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Control", "Section", "Title", "Status", "Applicable", "Justification"]
    for i, h in enumerate(headers):
        cell = soa_table.rows[0].cells[i]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)

    for r in rows:
        cells = soa_table.add_row().cells
        cells[0].text = r["control_id"]
        cells[1].text = r["section"]
        cells[2].text = r["title"]
        cells[3].text = r["status"]
        cells[4].text = "Yes" if r["applicable"] else "No"
        cells[5].text = r["justification"]
        for cell in cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)

    doc.add_paragraph("")
    doc.add_paragraph("Generated by Khestra — ISO/IEC 27001 Statement of Applicability.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


ISO_RISK_FRAMEWORKS = ("ISO 27001", "ISO27001", "iso27001")


def _risk_control_ids(risk: Dict[str, Any]) -> List[str]:
    ids: List[str] = []
    raw = risk.get("control_ids") or []
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except (json.JSONDecodeError, TypeError):
            raw = [raw] if raw else []
    if isinstance(raw, list):
        ids.extend(str(c) for c in raw if c)
    single = risk.get("control_id") or ""
    if single and str(single) not in ids:
        ids.append(str(single))
    return ids


def linked_risks(control_id: str) -> List[Dict[str, Any]]:
    """ISO-scoped risks that selected this control for treatment."""
    try:
        from risks.store import get_risks_by_control

        rows = get_risks_by_control(control_id) or []
    except Exception:
        return []
    out = []
    for r in rows:
        if (r.get("framework") or "") not in ISO_RISK_FRAMEWORKS:
            continue
        out.append({
            "id": r.get("id"),
            "title": r.get("title") or "",
            "status": r.get("status") or "",
            "treatment": r.get("treatment") or "",
            "inherent_score": r.get("inherent_score") or 0,
        })
    return out


def apply_treatment_to_soa(
    control_ids: List[str],
    risk_title: str = "",
    risk_id: str = "",
) -> List[str]:
    """Mark selected controls applicable. Does not set implemented."""
    note = (
        f"Selected for treatment of risk {risk_title} ({risk_id})."
        if risk_id else
        "Selected for treatment via the ISO 27001 risk register (clause 6.1.3)."
    )
    updated: List[str] = []
    for cid in control_ids:
        row = get_control(cid)
        if not row:
            continue
        fields: Dict[str, Any] = {}
        if not row["applicable"] or row["status"] == "excluded":
            fields["applicable"] = True
        if row["status"] == "excluded":
            fields["status"] = "not implemented"
        just = (row.get("justification") or "").strip()
        if not just:
            fields["justification"] = note
        elif row["status"] == "excluded" and risk_id and risk_id not in just:
            fields["justification"] = just + "\n" + note
        if fields:
            update_control(cid, fields)
            updated.append(cid)
    return updated


def sync_soa_from_risks() -> Dict[str, Any]:
    """Write risk-treatment control selections onto the SoA (clause 6.1.3)."""
    try:
        from risks.store import list_risks

        rows = list_risks(framework="ISO 27001") or []
    except Exception:
        return {"updated": [], "risks": 0}
    touched: List[str] = []
    for risk in rows:
        cids = _risk_control_ids(risk)
        if not cids:
            continue
        touched.extend(apply_treatment_to_soa(
            cids,
            risk_title=risk.get("title") or "",
            risk_id=risk.get("id") or "",
        ))
    return {"updated": sorted(set(touched)), "risks": len(rows)}
