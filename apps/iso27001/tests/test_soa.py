"""Tests for the ISO 27001 Statement of Applicability store + exports.

Run from repo root:
    apps/cmmc/.venv/bin/python -m pytest apps/iso27001/tests -q
"""

from __future__ import annotations

import os

import pytest

from controls import ALL_CONTROLS
from soa import (
    apply_treatment_to_soa,
    export_csv,
    export_docx,
    export_xlsx,
    get_control,
    init_store,
    linked_risks,
    list_soa,
    rollup,
    sync_soa_from_risks,
    update_control,
)


@pytest.fixture(autouse=True)
def _env(tmp_path):
    os.environ["RISKS_DB_PATH"] = str(tmp_path / "risks.db")
    init_store(str(tmp_path))
    yield
    for name in ("soa.db", "risks.db"):
        try:
            os.remove(str(tmp_path / name))
        except OSError:
            pass


class TestSoaStore:
    def test_seeds_all_controls(self):
        rows = list_soa()
        assert len(rows) == 116
        annex = [r for r in rows if r["control_id"].startswith("A.")]
        assert len(annex) == 93

    def test_default_status_not_implemented(self):
        row = get_control("A.5.1")
        assert row is not None
        assert row["status"] == "not implemented"
        assert row["applicable"] is True

    def test_update_status(self):
        updated = update_control("A.5.1", {"status": "implemented", "justification": "Policy live"})
        assert updated["status"] == "implemented"
        assert updated["justification"] == "Policy live"
        assert get_control("A.5.1")["status"] == "implemented"

    def test_update_rejects_invalid_status(self):
        with pytest.raises(ValueError, match="Invalid status"):
            update_control("A.5.1", {"status": "maybe"})

    def test_update_unknown_control_returns_none(self):
        assert update_control("A.99.99", {"status": "implemented"}) is None

    def test_exclusion_requires_justification(self):
        with pytest.raises(ValueError, match="justification"):
            update_control("A.7.3", {"status": "excluded", "applicable": False})
        updated = update_control("A.7.3", {
            "status": "excluded",
            "applicable": False,
            "justification": "Physical site not in scope",
        })
        assert updated["status"] == "excluded"
        assert updated["applicable"] is False
        assert "Physical site" in updated["justification"]

    def test_rollup_counts(self):
        update_control("A.5.1", {"status": "implemented"})
        update_control("A.5.2", {"status": "implemented"})
        update_control("A.5.3", {"status": "partially implemented"})
        update_control("A.7.3", {
            "status": "excluded",
            "applicable": False,
            "justification": "No on-premises physical infrastructure",
        })
        summary = rollup()
        assert summary["total"] == 116
        assert summary["counts"]["implemented"] == 2
        assert summary["counts"]["partially implemented"] == 1
        assert summary["counts"]["excluded"] == 1
        assert summary["applicable_total"] == 115

    def test_attributes_backfilled(self):
        row = get_control("A.5.1")
        assert isinstance(row["attributes"], dict)


class TestSoaExports:
    def test_export_csv(self):
        text = export_csv()
        assert text.startswith("\ufeff")
        assert "control_id,section,title,status" in text
        assert "A.5.1" in text

    def test_export_xlsx(self):
        from openpyxl import load_workbook
        from io import BytesIO

        data = export_xlsx()
        wb = load_workbook(BytesIO(data))
        ws = wb.active
        assert ws.title == "Statement of Applicability"
        assert ws.max_row == 117  # header + 116 controls
        assert ws.cell(1, 1).value == "Control"
        assert ws.cell(2, 1).value == "A.5.1"

    def test_export_docx(self):
        from docx import Document
        from io import BytesIO

        data = export_docx()
        doc = Document(BytesIO(data))
        assert len(doc.tables) == 2
        assert len(doc.tables[1].rows) == 117  # header + 116 controls
        assert "Statement of Applicability" in doc.paragraphs[0].text


class TestSoaRiskLink:
    def test_treatment_marks_applicable_not_implemented(self):
        update_control("A.8.32", {"status": "not implemented", "applicable": True})
        apply_treatment_to_soa(["A.8.32"], risk_title="Change risk", risk_id="risk1")
        row = get_control("A.8.32")
        assert row["applicable"] is True
        assert row["status"] == "not implemented"
        assert "Change risk" in row["justification"]
        assert "risk1" in row["justification"]

    def test_treatment_does_not_force_implemented(self):
        apply_treatment_to_soa(["A.5.1"], risk_title="Access risk", risk_id="r2")
        assert get_control("A.5.1")["status"] == "not implemented"

    def test_sync_from_risks_and_linked_risks(self):
        from risks.store import create_risk

        created = create_risk(
            title="Supply-chain compromise",
            framework="ISO 27001",
            control_ids=["A.5.21", "A.5.22"],
            treatment="mitigate",
        )
        result = sync_soa_from_risks()
        assert "A.5.21" in result["updated"]
        assert "Selected for treatment" in get_control("A.5.21")["justification"]
        linked = linked_risks("A.5.21")
        assert any(r["id"] == created["id"] for r in linked)
        assert linked[0]["title"] == "Supply-chain compromise"

    def test_sync_is_idempotent(self):
        from risks.store import create_risk

        create_risk(
            title="Repeat treatment",
            framework="ISO 27001",
            control_ids=["A.5.23"],
        )
        sync_soa_from_risks()
        first = get_control("A.5.23")["justification"]
        sync_soa_from_risks()
        assert get_control("A.5.23")["justification"] == first
