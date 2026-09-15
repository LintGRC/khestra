"""Tests for ISO 27001 auditor pack (clauses 9.2 / 9.3).

Run from repo root:
    apps/cmmc/.venv/bin/python -m pytest apps/iso27001/tests/test_audit_pack.py -q
"""

from __future__ import annotations

import os
import zipfile
from io import BytesIO

import pytest

from demo_seed import load_demo, seed_demo_audit_loop
from iso_audit_pack import (
    export_auditor_pack_zip,
    export_internal_audit_report_docx,
    export_management_review_minutes_docx,
    export_nc_capa_xlsx,
)
from soa import init_store as init_soa


@pytest.fixture(autouse=True)
def _env(tmp_path):
    os.environ["RISKS_DB_PATH"] = str(tmp_path / "risks.db")
    os.environ["MANAGEMENT_REVIEW_DB_PATH"] = str(tmp_path / "management_review.db")
    init_soa(str(tmp_path))
    from audit_center.store import init_store as init_audits
    from findings.store import init_store as init_findings
    from management_review.store import init_store as init_mr

    init_audits(str(tmp_path))
    init_findings(str(tmp_path))
    init_mr(str(tmp_path))
    yield


class TestAuditPackExports:
    def test_empty_safe_docx_xlsx_zip(self):
        assert export_internal_audit_report_docx().startswith(b"PK")
        assert export_management_review_minutes_docx().startswith(b"PK")
        assert export_nc_capa_xlsx().startswith(b"PK")
        z = zipfile.ZipFile(BytesIO(export_auditor_pack_zip()))
        names = z.namelist()
        assert "01-internal-audit-report.docx" in names
        assert "02-management-review-minutes.docx" in names
        assert "03-nc-capa-register.xlsx" in names

    def test_seed_populates_pack(self):
        from docx import Document
        from openpyxl import load_workbook

        created = seed_demo_audit_loop()
        assert created["audit"] is True
        assert created["finding"] is True
        assert created["review"] is True
        again = seed_demo_audit_loop()
        assert again == {"audit": False, "finding": False, "review": False}

        audit_doc = Document(BytesIO(export_internal_audit_report_docx()))
        audit_text = "\n".join(p.text for p in audit_doc.paragraphs)
        assert "ISMS internal audit 2026-H1" in audit_text
        assert "Unpatched production image" in audit_text

        mr_doc = Document(BytesIO(export_management_review_minutes_docx()))
        mr_text = "\n".join(p.text for p in mr_doc.paragraphs)
        assert "Q1 2026 ISMS management review" in mr_text
        assert "clause 9.3" in mr_text.lower() or "9.3" in mr_text

        wb = load_workbook(BytesIO(export_nc_capa_xlsx()))
        ws = wb.active
        titles = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
        assert any(t and "Unpatched" in str(t) for t in titles)

    def test_load_demo_includes_audit_loop(self):
        out = load_demo()
        assert out["audit_loop"]["audit"] is True
        assert out["audit_loop"]["review"] is True
