"""Tests for the ISO 27001 risk assessment report + demo seed.

Run from repo root:
    apps/cmmc/.venv/bin/python -m pytest apps/iso27001/tests/test_risk_and_demo.py -q
"""

from __future__ import annotations

import os

import pytest

from demo_seed import load_demo, seed_demo_risks, seed_demo_soa
from iso_risk_report import export_risk_report_docx, export_risk_report_xlsx, risk_rows
from soa import init_store, rollup


@pytest.fixture(autouse=True)
def _env(tmp_path):
    os.environ["RISKS_DB_PATH"] = str(tmp_path / "risks.db")
    init_store(str(tmp_path))
    yield
    for name in ("risks.db", "soa.db"):
        try:
            os.remove(str(tmp_path / name))
        except OSError:
            pass


class TestDemoSeed:
    def test_soa_seed_distribution(self):
        seed_demo_soa()
        summary = rollup()
        assert summary["total"] == 116
        assert summary["counts"]["excluded"] == 7
        assert summary["counts"]["implemented"] >= 70

    def test_soa_seed_is_deterministic(self):
        seed_demo_soa()
        first = rollup()["counts"]
        seed_demo_soa()
        second = rollup()["counts"]
        assert first == second

    def test_risk_seed_creates_and_is_idempotent(self):
        created = seed_demo_risks()
        assert len(created) == 3
        again = seed_demo_risks()
        assert again == []

    def test_load_demo_returns_rollup(self):
        out = load_demo()
        assert out["status"] == "ok"
        assert len(out["risks_created"]) == 3
        assert out["soa_rollup"]["total"] == 116
        from soa import get_control

        treated = get_control("A.5.15")
        assert treated["applicable"] is True
        assert "Selected for treatment" in treated["justification"]


class TestRiskReport:
    def test_risk_rows_empty_safe(self):
        assert risk_rows() == []

    def test_risk_rows_with_seed(self):
        seed_demo_risks()
        rows = risk_rows()
        assert len(rows) == 3
        row = rows[0]
        assert row["title"]
        assert row["inherent_score"] == row["likelihood"] * row["impact"]
        assert row["treatment_label"]

    def test_docx_export_empty_safe(self):
        data = export_risk_report_docx()
        assert data.startswith(b"PK")  # valid docx zip

    def test_docx_export_with_risks(self):
        from docx import Document
        from io import BytesIO

        seed_demo_risks()
        doc = Document(BytesIO(export_risk_report_docx()))
        assert len(doc.tables) >= 3  # criteria + summary + register
        assert "Risk Assessment Report" in doc.paragraphs[0].text

    def test_xlsx_export_with_risks(self):
        from openpyxl import load_workbook
        from io import BytesIO

        seed_demo_risks()
        wb = load_workbook(BytesIO(export_risk_report_xlsx()))
        assert "Risk Register" in wb.sheetnames
        assert "Summary" in wb.sheetnames
        ws = wb["Risk Register"]
        assert ws.max_row == 4  # header + 3 risks
        assert ws.cell(2, 1).value  # risk id present
