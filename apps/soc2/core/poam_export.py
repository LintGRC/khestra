"""SOC 2 Exceptions export."""

from __future__ import annotations

import csv
import io
from typing import Any, Dict, List

from soc2_catalog import SOC2_CONTROLS


POAM_COLUMNS = [
    "Control ID",
    "Control Name",
    "Status",
    "Risk Level",
    "Weakness Description",
    "Compensating Controls",
    "Risk Acceptance",
    "Remediation Plan",
    "Owner",
    "Target Date",
    "Expiry Date",
    "Created At",
    "Updated At",
]


def build_poam_rows(
    answers: Dict[str, Any],
    exceptions: List[Dict[str, Any]],
) -> List[Dict[str, str]]:
    """Build POA&M rows from open gaps and exceptions."""

    rows: List[Dict[str, str]] = []

    # Include open gap controls (NOT MET, NOT STARTED)
    for cid, meta in SOC2_CONTROLS.items():
        ans = answers.get(cid, {})
        status = ans.get("status", "NOT STARTED")
        if status in ("NOT MET", "NOT STARTED"):
            rows.append({
                "Control ID": cid,
                "Control Name": meta["title"],
                "Status": status,
                "Risk Level": "",
                "Weakness Description": ans.get("implementation_narrative") or f"Control not yet implemented",
                "Compensating Controls": "",
                "Risk Acceptance": "",
                "Remediation Plan": ans.get("remediation_plan") or "",
                "Owner": ans.get("owner") or "",
                "Target Date": ans.get("target_date") or "",
                "Expiry Date": "",
                "Created At": "",
                "Updated At": "",
            })

    # Include active exceptions
    for exc in exceptions:
        if exc.get("status") in ("pending_approval", "approved"):
            cid = exc.get("control_id", "")
            meta = SOC2_CONTROLS.get(cid, {})
            rows.append({
                "Control ID": cid,
                "Control Name": meta.get("title", ""),
                "Status": exc.get("status", ""),
                "Risk Level": exc.get("risk_level", ""),
                "Weakness Description": exc.get("description", ""),
                "Compensating Controls": exc.get("compensating_controls", ""),
                "Risk Acceptance": exc.get("risk_acceptance", ""),
                "Remediation Plan": "",
                "Owner": exc.get("created_by", ""),
                "Target Date": "",
                "Expiry Date": exc.get("expiry_date", ""),
                "Created At": exc.get("created_at", ""),
                "Updated At": exc.get("updated_at", ""),
            })

    return rows


def export_poam_csv(
    answers: Dict[str, Any],
    exceptions: List[Dict[str, Any]],
) -> bytes:
    """Export POA&M as CSV bytes."""

    rows = build_poam_rows(answers, exceptions)
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=POAM_COLUMNS)
    writer.writeheader()
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def export_poam_docx(
    org_name: str,
    answers: Dict[str, Any],
    exceptions: List[Dict[str, Any]],
) -> bytes:
    """Export POA&M as DOCX bytes (table format)."""

    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    # Title
    title = doc.add_heading(f"POA&M — {org_name}", level=1)
    for run in title.runs:
        from docx.shared import RGBColor
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

    rows = build_poam_rows(answers, exceptions)

    if not rows:
        doc.add_paragraph("No open gaps or active POA&M items.")
    else:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        headers = ["Control", "Status", "Risk", "Description", "Owner", "Target"]
        for i, h in enumerate(headers):
            hdr[i].text = h
            for run in hdr[i].paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)

        for row in rows:
            cells = table.add_row().cells
            vals = [
                row["Control ID"],
                row["Status"],
                row["Risk Level"],
                row["Weakness Description"][:100],
                row["Owner"],
                row["Target Date"],
            ]
            for i, v in enumerate(vals):
                cells[i].text = v
                for para in cells[i].paragraphs:
                    para.style.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
