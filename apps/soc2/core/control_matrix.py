"""Control Matrix export — criteria × evidence × coverage grid as Excel."""

from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def generate_control_matrix(ws: Dict[str, Any]) -> bytes:
    from readiness import compute_dashboard
    from workspace_service import list_control_summaries

    dash = compute_dashboard(ws)
    controls = list_control_summaries(ws)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Control Matrix"

    bold = Font(bold=True)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    wrap = Alignment(wrap_text=True, vertical="top")
    thin = Side(style="thin")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)

    org = ws.get("org_name") or "Organization"
    ws_summary.merge_cells("A1:L1")
    c = ws_summary["A1"]
    c.value = f"Control Matrix — {org}"
    c.font = Font(bold=True, size=14)
    c.alignment = Alignment(horizontal="left")

    ws_summary.merge_cells("A2:L2")
    c = ws_summary["A2"]
    c.value = f"Total controls: {dash.get('controls_total', 0)} | Met: {dash.get('controls_met', 0)} | Readiness: {dash.get('readiness_pct', 0):.0f}% | Evidence coverage: {dash.get('evidence_coverage_pct', 0):.0f}%"
    c.font = Font(size=10, italic=True)

    headers = [
        "Control ID", "Control Name", "Category", "Status",
        "Has Evidence", "Evidence Count", "Auto Evidence",
        "PoF Coverage %", "Has Narrative", "Owner",
        "Target Date", "Operating Effectiveness",
    ]
    row = 4
    for col_idx, h in enumerate(headers, 1):
        cell = ws_summary.cell(row=row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border

    category_fills = {
        "Security": PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid"),
        "Availability": PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
        "Confidentiality": PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid"),
        "Processing Integrity": PatternFill(start_color="F3E5F5", end_color="F3E5F5", fill_type="solid"),
        "Privacy": PatternFill(start_color="FCE4EC", end_color="FCE4EC", fill_type="solid"),
    }

    status_fill = PatternFill(start_color="D32F2F", end_color="D32F2F", fill_type="solid")
    status_font = Font(bold=True, color="FFFFFF", size=9)

    for i, ctrl in enumerate(controls, row + 1):
        cat = ctrl.get("category", "").title()
        cat_fill = category_fills.get(cat)
        ev_count = ctrl.get("evidence_count", 0)
        pof_pct = ctrl.get("pof_coverage_pct", 0) or 0
        has_evidence = "Y" if ev_count > 0 else "N"
        has_auto = "Y" if ctrl.get("has_auto_evidence") else "N"
        has_narrative = "Y" if ctrl.get("has_narrative") else "N"
        status = ctrl.get("status", "NOT STARTED")

        vals = [
            ctrl.get("code", ""),
            ctrl.get("name", ""),
            cat,
            status,
            has_evidence,
            ev_count,
            has_auto,
            pof_pct,
            has_narrative,
            ctrl.get("owner", ""),
            ctrl.get("target_date", ""),
            ctrl.get("operating_status", "NOT TESTED"),
        ]
        for col_idx, v in enumerate(vals, 1):
            cell = ws_summary.cell(row=i, column=col_idx, value=v)
            cell.alignment = wrap
            cell.border = border
            if cat_fill:
                cell.fill = cat_fill

    # Column widths
    widths = [12, 30, 16, 14, 12, 14, 12, 14, 12, 20, 14, 20]
    for idx, w in enumerate(widths, 1):
        ws_summary.column_dimensions[get_column_letter(idx)].width = w

    # Freeze header row
    ws_summary.freeze_panes = "A5"

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
