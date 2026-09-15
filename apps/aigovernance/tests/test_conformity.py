"""Tests for conformity framework completeness + EU artifact generators.

Run from repo root:
    apps/aigovernance/.venv/bin/python -m pytest apps/aigovernance/tests/test_conformity.py -q
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from core.iso_42001_official import ANNEX_A_CONTROL_TOPICS as ISO_CONTROL_TOPICS  # noqa: E402
from core.ai_controls_catalog import AI_GOV_FRAMEWORK  # noqa: E402
from server.conformity_routes import FRAMEWORKS, applicable_articles  # noqa: E402
from server.conformity_routes import (  # noqa: E402
    EU_DEFAULT_APPLIES_FROM,
    _eu_article_in_force,
    eu_applicable_from,
)


def test_eu_ai_act_applicability_dates():
    """EU AI Act phases in article obligations per the 2026 Omnibus dates."""
    # Prohibitions & GPAI model rules long in force.
    assert eu_applicable_from({"id": "5"}) == "2025-02-02"
    assert eu_applicable_from({"id": "51"}) == "2025-08-02"
    assert eu_applicable_from({"id": "56"}) == "2025-08-02"
    # Transparency + enforcement powers.
    assert eu_applicable_from({"id": "50"}) == "2026-08-02"
    # High-risk (Annex III standalone) default.
    assert eu_applicable_from({"id": "43"}) == EU_DEFAULT_APPLIES_FROM == "2027-12-02"
    assert eu_applicable_from({"id": "8"}) == "2027-12-02"
    # As of today (2026), Annex III high-risk obligations are not yet in force.
    assert not _eu_article_in_force({"id": "43"}, "2026-08-17")
    assert _eu_article_in_force({"id": "5"}, "2026-08-17")


def test_eu_ai_act_article_count():
    fw = FRAMEWORKS["eu_ai_act"]
    assert len(fw["articles"]) >= 40
    refs = {a["control_id"] for a in fw["articles"]}
    for expected in ("EU-8", "EU-46", "EU-47", "EU-57", "EU-60", "EU-73"):
        assert expected in refs


def test_eu_articles_have_tiers_and_ids():
    for a in FRAMEWORKS["eu_ai_act"]["articles"]:
        assert a["id"] and a["control_id"] and a["tiers"] and a["ref"]


def test_nist_ai_rmf_full_category_coverage():
    fw = FRAMEWORKS["nist_ai_rmf"]
    cats = {a["control_id"] for a in fw["articles"]}
    # RMF 1.0: 72 subcategories — GOVERN 19, MAP 18, MEASURE 22, MANAGE 13
    assert len(fw["articles"]) == 72
    for prefix, count in [("NIST-GOVERN-", 19), ("NIST-MAP-", 18), ("NIST-MEASURE-", 22), ("NIST-MANAGE-", 13)]:
        got = {c for c in cats if c.startswith(prefix)}
        assert len(got) == count, f"{prefix} expected {count}, got {len(got)}"


def test_nist_ai_rmf_articles_match_catalog():
    """Conformity NIST articles must be exactly the catalog's NIST set (single source)."""
    fw = FRAMEWORKS["nist_ai_rmf"]
    cids = {a["control_id"] for a in fw["articles"]}
    cat = {k for k in AI_GOV_FRAMEWORK if k.startswith("NIST-")}
    assert cids == cat, f"conformity/catalog mismatch: extra={cids - cat}, missing={cat - cids}"


def test_eu_ai_act_articles_match_catalog():
    """Conformity EU articles must be exactly the catalog's EU control set (single source)."""
    fw = FRAMEWORKS["eu_ai_act"]
    cids = {a["control_id"] for a in fw["articles"]}
    cat = {k for k in AI_GOV_FRAMEWORK if k.startswith("EU-")}
    assert cids == cat, f"conformity/catalog mismatch: extra={cids - cat}, missing={cat - cids}"


def test_owasp_agentic_articles_match_catalog():
    """Conformity OWASP articles must be exactly the catalog's OWASP set (single source)."""
    fw = FRAMEWORKS["owasp_agentic"]
    cids = {a["control_id"] for a in fw["articles"]}
    cat = {k for k in AI_GOV_FRAMEWORK if k.startswith("OWASP-") and not k.startswith("OWASP-LLM-")}
    assert cids == cat, f"conformity/catalog mismatch: extra={cids - cat}, missing={cat - cids}"
    assert len(fw["articles"]) == 10


def test_owasp_llm_articles_match_catalog():
    """Conformity OWASP LLM articles must be exactly the catalog's OWASP LLM set."""
    fw = FRAMEWORKS["owasp_llm"]
    cids = {a["control_id"] for a in fw["articles"]}
    cat = {k for k in AI_GOV_FRAMEWORK if k.startswith("OWASP-LLM-")}
    assert cids == cat, f"conformity/catalog mismatch: extra={cids - cat}, missing={cat - cids}"
    assert len(fw["articles"]) == 10


def test_iso_42001_clauses_and_annex():
    fw = FRAMEWORKS["iso_42001"]
    clause_refs = [a for a in fw["articles"] if a["ref"].startswith("Clause")]
    annex = [a for a in fw["articles"] if not a["ref"].startswith("Clause")]
    assert len(clause_refs) == 23
    assert len(annex) == 38
    expected_refs = set()
    for topics in ISO_CONTROL_TOPICS.values():
        expected_refs.update(topics)
    assert {a["ref"] for a in annex} == expected_refs
    # control articles must map to real catalog controls
    assert all(a["control_id"] in AI_GOV_FRAMEWORK for a in annex)


SYSTEM = {
    "id": "test-sys",
    "system_id": "test-sys",
    "name": "Test Classifier",
    "version": "1.0.0",
    "vendor": "Test Org",
    "risk_classification": "high",
    "purpose": "Automated decision support",
    "description": "A test system",
    "input_data_sources": ["employee records"],
    "output_destinations": ["HR portal"],
    "processing_location": "EU",
}


class TestArtifactGenerators:
    def test_eu_doc_generates_docx(self):
        from server.eu_artifacts import export_eu_declaration_of_conformity

        data = export_eu_declaration_of_conformity(SYSTEM)
        assert data.startswith(b"PK")

    def test_tech_doc_generates_docx(self):
        from server.eu_artifacts import export_technical_documentation

        data = export_technical_documentation(SYSTEM)
        assert data.startswith(b"PK")

    def test_ai_soa_generates_docx(self):
        from server.eu_artifacts import export_ai_soa

        data = export_ai_soa(SYSTEM, "iso_42001")
        assert data.startswith(b"PK")

    def test_ai_soa_rejects_unknown_framework(self):
        from server.eu_artifacts import export_ai_soa

        with pytest.raises(ValueError, match="Unknown framework"):
            export_ai_soa(SYSTEM, "nope")

    def test_ai_soa_applicable_no_when_na(self):
        from docx import Document
        from io import BytesIO
        from server.eu_artifacts import _render_soa_table

        doc = Document()
        _render_soa_table(doc, [{
            "ref": "A.2.2",
            "title": "AI policy",
            "status": "na",
            "notes": "Not in scope — no model training.",
        }])
        row = doc.tables[0].rows[1]
        assert row.cells[3].text == "No"
        assert "Not in scope" in row.cells[4].text

    def test_ai_soa_applicable_yes_when_assessed(self):
        from docx import Document
        from server.eu_artifacts import _render_soa_table

        doc = Document()
        _render_soa_table(doc, [{
            "ref": "A.2.3",
            "title": "AI roles",
            "status": "compliant",
            "notes": "",
        }])
        assert doc.tables[0].rows[1].cells[3].text == "Yes"

    def test_dossier_pack_zip(self):
        import zipfile
        from io import BytesIO
        from server.eu_artifacts import export_dossier_pack

        data = export_dossier_pack(SYSTEM)
        assert data.startswith(b"PK")
        names = zipfile.ZipFile(BytesIO(data)).namelist()
        assert "01-eu-declaration-of-conformity.docx" in names
        assert "02-technical-documentation.docx" in names
        assert "03-ai-statement-of-applicability.docx" in names
        assert "04-model-card.md" in names
        assert "README.txt" in names


class TestClassifyScopedQueue:
    def test_eu_deployer_hides_provider_articles(self):
        # as_of after all high-risk obligations are in force, to test pure scoping.
        as_of = "2028-01-01"
        provider_sys = {**SYSTEM, "risk_classification": "high", "ai_profile": "provider"}
        deployer_sys = {**SYSTEM, "risk_classification": "high", "ai_profile": "deployer"}
        provider_ids = {a["id"] for a in applicable_articles(provider_sys, "eu_ai_act", as_of)}
        deployer_ids = {a["id"] for a in applicable_articles(deployer_sys, "eu_ai_act", as_of)}
        assert "16" in provider_ids  # Art. 16 provider QMS
        assert "16" not in deployer_ids
        assert "26" in deployer_ids  # Art. 26 deployer
        assert len(deployer_ids) < len(provider_ids)

    def test_owasp_agentic_only_for_agentic_profile(self):
        provider = {**SYSTEM, "risk_classification": "high", "ai_profile": "provider"}
        agentic = {**SYSTEM, "risk_classification": "high", "ai_profile": "agentic"}
        assert applicable_articles(provider, "owasp_agentic") == []
        queued = applicable_articles(agentic, "owasp_agentic")
        assert len(queued) == 10

    def test_nist_minimal_is_govern_only(self):
        minimal = {**SYSTEM, "risk_classification": "minimal", "ai_profile": "provider"}
        high = {**SYSTEM, "risk_classification": "high", "ai_profile": "provider"}
        min_ids = {a["control_id"] for a in applicable_articles(minimal, "nist_ai_rmf")}
        high_ids = {a["control_id"] for a in applicable_articles(high, "nist_ai_rmf")}
        assert all(c.startswith("NIST-GOVERN-") for c in min_ids)
        assert any(c.startswith("NIST-MEASURE-") for c in high_ids)
        assert len(min_ids) < len(high_ids)

    def test_iso_minimal_hides_operation_clauses(self):
        minimal = {**SYSTEM, "risk_classification": "minimal", "ai_profile": "provider"}
        high = {**SYSTEM, "risk_classification": "high", "ai_profile": "provider"}
        min_cats = {a["category"] for a in applicable_articles(minimal, "iso_42001")}
        high_n = len(applicable_articles(high, "iso_42001"))
        assert "operation" not in min_cats
        assert high_n > len(applicable_articles(minimal, "iso_42001"))

    def test_gap_analysis_skips_articles_outside_queue(self):
        from server.conformity_routes import gap_analysis_rows

        systems = {
            "a": {**SYSTEM, "risk_classification": "minimal", "ai_profile": "provider", "controls": []},
        }
        nist = gap_analysis_rows(systems, "nist_ai_rmf")
        ids = {r["control_id"] for r in nist}
        assert any(i.startswith("GOVERN") for i in ids)
        assert not any(i.startswith("MEASURE") for i in ids)
        assert not any(i.startswith("MAP") for i in ids)
        agentic = gap_analysis_rows(systems, "owasp_agentic")
        assert agentic == []
