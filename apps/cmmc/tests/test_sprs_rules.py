"""Pure-function tests for SPRS scoring, POA&M eligibility, and what-if rules.

These are the compliance math — if poam_eligibility() or simulate_sprs() is
wrong, the tool's SPRS output is wrong. Run from repo root:
    apps/cmmc/.venv/bin/python -m pytest apps/cmmc/tests/test_sprs_rules.py -q
"""

import os
import sys
from pathlib import Path

import pytest

APP = Path(__file__).resolve().parents[1]
CORE = APP / "core"
if str(CORE) not in sys.path:
    sys.path.insert(0, str(CORE))
if str(APP.parents[1] / "packages" / "evidence") not in sys.path:
    sys.path.insert(0, str(APP.parents[1] / "packages" / "evidence"))

from controls import CMMC_FRAMEWORK  # noqa: E402
from poam_eligibility import poam_eligibility, unmet_controls  # noqa: E402
from poam_export import build_poam_dataframe, poam_weakness_ids  # noqa: E402
from readiness import build_control_ledger, simulate_sprs  # noqa: E402
from sprs_engine import calculate_detailed_sprs  # noqa: E402

SCOPED = list(CMMC_FRAMEWORK.keys())


def all_met(**overrides):
    answers = {c: {"status": "MET"} for c in SCOPED}
    for cid, ans in overrides.items():
        answers[cid] = ans
    return answers


# ── SPRS scoring honesty rules ───────────────────────────────────────────────


def test_all_met_scores_110():
    assert calculate_detailed_sprs(all_met(), SCOPED)["final_score"] == 110


def test_unanswered_controls_are_deducted_and_surfaced():
    result = calculate_detailed_sprs({}, SCOPED)
    assert result["final_score"] == -203  # floored
    assert result["unanswered_count"] == 110
    assert len(result["unanswered_ids"]) == 110
    assert result["answered_count"] == 0


def test_unjustified_na_scores_as_not_met():
    answers = all_met()
    answers["AC.L2-3.1.1"] = {"status": "NOT APPLICABLE"}  # 5-point, no justification
    answers["AC.L2-3.1.2"] = {"status": "NOT APPLICABLE", "justification": "no CUI flow"}
    result = calculate_detailed_sprs(answers, SCOPED)
    assert result["final_score"] == 105  # only the unjustified 5-point N/A deducted
    assert "AC.L2-3.1.1" in result["critical_gaps"]


def test_variable_mfa_and_fips_deductions():
    answers = all_met()
    answers["IA.L2-3.5.3"] = {"status": "PARTIALLY MET"}  # 3
    answers["SC.L2-3.13.11"] = {"status": "NOT MET"}  # 5
    result = calculate_detailed_sprs(answers, SCOPED)
    assert result["final_score"] == 102
    assert result["mfa_deduction"] == 3
    assert result["fips_deduction"] == 5


# ── POA&M eligibility (32 CFR 170.21) ────────────────────────────────────────


def test_one_point_gap_is_eligible():
    answers = all_met()
    answers["AC.L2-3.1.3"] = {"status": "NOT MET"}  # 1-point
    e = poam_eligibility(answers, SCOPED)
    assert e["eligible"] is True
    assert e["score"] == 109
    assert e["eligible_ids"] == ["AC.L2-3.1.3"]
    assert e["blocking_ids"] == []


def test_five_point_gap_blocks_eligibility():
    answers = all_met()
    answers["AC.L2-3.1.1"] = {"status": "NOT MET"}  # 5-point
    e = poam_eligibility(answers, SCOPED)
    assert e["eligible"] is False
    assert "AC.L2-3.1.1" in e["blocking_ids"]


def test_score_below_88_blocks_eligibility():
    answers = all_met()
    # knock out enough 5-point controls to fall below 88
    for cid in ("AC.L2-3.1.1", "AC.L2-3.1.2", "AC.L2-3.1.12", "AC.L2-3.1.13", "AC.L2-3.1.16"):
        answers[cid] = {"status": "NOT MET"}
    e = poam_eligibility(answers, SCOPED)
    assert e["score"] < 88
    assert e["eligible"] is False


def test_fips_partial_rides_poam_at_3_points():
    answers = all_met()
    answers["SC.L2-3.13.11"] = {"status": "PARTIALLY MET"}  # 3-point exception
    e = poam_eligibility(answers, SCOPED)
    assert e["eligible"] is True
    assert "SC.L2-3.13.11" in e["eligible_ids"]


def test_fips_full_gap_blocks_eligibility():
    answers = all_met()
    answers["SC.L2-3.13.11"] = {"status": "NOT MET"}  # 5-point
    e = poam_eligibility(answers, SCOPED)
    assert e["eligible"] is False
    assert "SC.L2-3.13.11" in e["blocking_ids"]


def test_mfa_gap_blocks_eligibility():
    answers = all_met()
    answers["IA.L2-3.5.3"] = {"status": "PARTIALLY MET"}
    e = poam_eligibility(answers, SCOPED)
    assert e["eligible"] is False
    assert "IA.L2-3.5.3" in e["blocking_ids"]


def test_na_controls_not_unmet_when_justified():
    answers = all_met()
    answers["AC.L2-3.1.2"] = {"status": "NOT APPLICABLE", "justification": "no CUI flow"}
    assert "AC.L2-3.1.2" not in unmet_controls(answers, SCOPED)


def test_unjustified_na_is_unmet():
    answers = all_met()
    answers["AC.L2-3.1.2"] = {"status": "NOT APPLICABLE"}
    assert "AC.L2-3.1.2" in unmet_controls(answers, SCOPED)


# ── What-if simulator ────────────────────────────────────────────────────────


def test_simulate_delta_negative_for_gap():
    sim = simulate_sprs(all_met(), SCOPED, [{"control_id": "AC.L2-3.1.1", "status": "NOT MET"}])
    assert sim["current_score"] == 110
    assert sim["projected_score"] == 105
    assert sim["delta"] == -5
    assert sim["applied_changes"] == ["AC.L2-3.1.1"]


def test_simulate_ignores_unknown_controls():
    sim = simulate_sprs(all_met(), SCOPED, [{"control_id": "NOPE.L2-9.9.9", "status": "NOT MET"}])
    assert sim["delta"] == 0
    assert sim["applied_changes"] == []


def test_simulate_improves_score_when_closing_gap():
    answers = all_met()
    answers["AC.L2-3.1.1"] = {"status": "NOT MET"}
    sim = simulate_sprs(answers, SCOPED, [{"control_id": "AC.L2-3.1.1", "status": "MET"}])
    assert sim["delta"] == 5


# ── Control ledger ───────────────────────────────────────────────────────────


def test_ledger_rows_and_poam_flag():
    answers = all_met()
    answers["AC.L2-3.1.3"] = {"status": "NOT MET"}
    ledger = build_control_ledger(answers, SCOPED)
    assert len(ledger) == 110
    row = next(r for r in ledger if r["control_id"] == "AC.L2-3.1.3")
    assert row["weight"] == 1
    assert row["deduction"] == 1
    assert row["poam_eligible"] is True
    assert row["assessment_ready"] is False  # NOT MET -> not assessment-ready
    met_row = next(r for r in ledger if r["control_id"] == "AC.L2-3.1.1")
    assert met_row["status"] == "MET"
    assert met_row["assessment_ready"] is False  # MET but no evidence attached


# ── POA&M weakness IDs + SSP cross-reference ─────────────────────────────────


def test_poam_weakness_ids_match_export_rows():
    answers = all_met()
    answers["AC.L2-3.1.1"] = {"status": "NOT MET"}  # 5-point first
    answers["AC.L2-3.1.3"] = {"status": "PARTIALLY MET"}
    answers["AC.L2-3.1.5"] = {"status": "NOT APPLICABLE"}  # unjustified -> gap
    ids = poam_weakness_ids(answers, SCOPED)
    df = build_poam_dataframe(answers, SCOPED)
    rows = df[["Control ACID", "Weakness ID"]].to_dict("records")
    assert [r["Weakness ID"] for r in rows] == [ids[r["Control ACID"]] for r in rows]
    assert rows[0]["Control ACID"] == "AC.L2-3.1.1"  # 5-point sorted first
    assert ids["AC.L2-3.1.1"] == "WK-0001"
    assert "AC.L2-3.1.5" in ids  # unjustified N/A appears on the POA&M
    assert "AC.L2-3.1.6" not in ids if "AC.L2-3.1.6" in SCOPED else True


def test_ssp_cross_reference_text():
    """Non-MET controls render 'remediation tracked in POA&M Item WK-NNNN'."""
    from docx import Document
    from ssp.sections.control_family import add_control_family_section

    answers = all_met()
    answers["AC.L2-3.1.1"] = {"status": "NOT MET"}
    answers["AC.L2-3.1.3"] = {"status": "NOT MET"}
    ids = poam_weakness_ids(answers, SCOPED)

    doc = Document()
    controls = [(cid, answers[cid]) for cid in ("AC.L2-3.1.1", "AC.L2-3.1.3")]
    add_control_family_section(
        doc, 1, "Access Control", controls,
        page_break_after=False, poam_ids=ids,
    )
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "POA&M Item WK-0001" in text
    assert "POA&M Item WK-0002" in text
    assert "Control not yet implemented" in text


def test_ssp_cross_reference_with_narrative():
    """Gap controls WITH a narrative still get the POA&M linkage line."""
    from docx import Document
    from ssp.sections.control_family import add_control_family_section

    answers = all_met()
    answers["AC.L2-3.1.1"] = {"status": "NOT MET", "implementation_narrative": "Partial rollout in progress"}
    ids = poam_weakness_ids(answers, SCOPED)

    doc = Document()
    add_control_family_section(
        doc, 1, "Access Control", [("AC.L2-3.1.1", answers["AC.L2-3.1.1"])],
        page_break_after=False, poam_ids=ids,
    )
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Partial rollout in progress" in text
    assert "POA&M Item WK-0001" in text


def test_ssp_met_controls_have_no_poam_link():
    """MET controls must not carry a POA&M cross-reference."""
    from docx import Document
    from ssp.sections.control_family import add_control_family_section

    answers = all_met()  # all MET
    ids = poam_weakness_ids(answers, SCOPED)
    assert ids == {}

    doc = Document()
    add_control_family_section(
        doc, 1, "Access Control", [("AC.L2-3.1.1", answers["AC.L2-3.1.1"])],
        page_break_after=False, poam_ids=ids,
    )
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "POA&M Item" not in text


def test_template_cell_references_poam_item():
    """Official CUI template cell narrative references the specific POA&M item."""
    from ssp.template_filler import _narrative_for_cell

    # gap without narrative -> explicit item ref
    col, text = _narrative_for_cell("NOT MET", "", "AC.L2-3.1.1", "WK-0001")
    assert col == 1
    assert "POA&M Item WK-0001" in text

    # gap with narrative -> ref appended
    col, text = _narrative_for_cell("PARTIALLY MET", "Partial rollout", "AC.L2-3.1.3", "WK-0003")
    assert col == 1
    assert "Partial rollout" in text and "POA&M Item WK-0003" in text

    # no poam map -> generic fallback (no crash)
    col, text = _narrative_for_cell("NOT MET", "", "AC.L2-3.1.1", None)
    assert col == 1
    assert "the POA&M" in text

    # MET -> no ref
    col, text = _narrative_for_cell("MET", "Fully implemented", "AC.L2-3.1.1", "WK-0001")
    assert col == 0
    assert "POA&M" not in text

    # justified N/A -> no ref
    col, text = _narrative_for_cell("NOT APPLICABLE", "", "AC.L2-3.1.6", None)
    assert col == 2
    assert "POA&M" not in text


def test_oscal_poam_export_shape():
    """OSCAL POA&M export: one poam-item per gap, linked to weakness IDs."""
    import json
    from workspace_io_platform import export_oscal_poam_json

    answers = all_met()
    answers["AC.L2-3.1.1"] = {"status": "NOT MET"}
    answers["AC.L2-3.1.3"] = {"status": "PARTIALLY MET", "remediation_plan": "Enable MFA", "target_date": "2026-12-31"}
    out = json.loads(export_oscal_poam_json(answers, SCOPED, "TestOrg"))
    poam = out["plan-of-action-and-milestones"]
    assert poam["metadata"]["oscal-version"] == "1.0.4"
    items = poam["poam-items"]
    assert len(items) == 2
    first = items[0]
    props = {p["name"]: p["value"] for p in first["props"]}
    assert props["control-id"] == "AC.L2-3.1.1"
    assert props["weakness-id"] == "WK-0001"
    assert first["related-requirements"] == [{"requirement-id": "AC.L2-3.1.1"}]
    second_props = {p["name"]: p["value"] for p in items[1]["props"]}
    assert second_props["control-id"] == "AC.L2-3.1.3"
    assert items[1]["remediation"]["target-date"] == "2026-12-31"


def test_ssp_falls_back_without_poam_map():
    from docx import Document
    from ssp.sections.control_family import add_control_family_section

    answers = all_met()
    answers["AC.L2-3.1.1"] = {"status": "NOT MET"}
    doc = Document()
    add_control_family_section(
        doc, 1, "Access Control", [("AC.L2-3.1.1", answers["AC.L2-3.1.1"])],
        page_break_after=False, poam_ids=None,
    )
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Control not yet implemented." in text
    assert "POA&M Item" not in text


def test_executive_report_builds_all_sections():
    """Board report PDF: all five sections render, resource ask included."""
    from pdf_reports import build_executive_report

    answers = all_met()
    answers["AC.L2-3.1.1"] = {
        "status": "NOT MET",
        "estimated_cost": "15000",
        "target_date": "2025-01-01",
    }
    answers["AC.L2-3.1.3"] = {
        "status": "PARTIALLY MET",
        "estimated_cost": "2000",
        "target_date": "2026-12-31",
    }
    sprs = calculate_detailed_sprs(answers, SCOPED)
    eligibility = poam_eligibility(answers, SCOPED)
    ws = {
        "org_profile": {
            "org_name": "TestOrg",
            "system_name": "Test System",
            "board_resource_ask": "Need $15k for MDM",
        },
        "answers": answers,
        "scoped_controls": SCOPED,
        "sprs_history": [
            {"timestamp": "2026-07-01T00:00:00", "score": 80},
            {"timestamp": "2026-08-01T00:00:00", "score": 85},
        ],
    }
    pdf_bytes = build_executive_report(ws, sprs, eligibility, [])
    assert pdf_bytes.startswith(b"%PDF")
    text = _pdf_text(pdf_bytes)
    assert "Executive Board Report" in text
    assert "SPRS Score" in text
    assert "Posture Over Time" in text
    assert "Top Risks" in text
    assert "POA&M Health" in text
    assert "Resource Ask" in text
    assert "MDM" in text
    assert "Overdue Items" in text


def test_executive_report_skips_resource_ask_when_empty():
    from pdf_reports import build_executive_report

    answers = all_met()
    sprs = calculate_detailed_sprs(answers, SCOPED)
    ws = {
        "org_profile": {"org_name": "TestOrg", "board_resource_ask": ""},
        "answers": answers,
        "scoped_controls": SCOPED,
        "sprs_history": [],
    }
    pdf_bytes = build_executive_report(ws, sprs, {}, [])
    text = _pdf_text(pdf_bytes)
    assert "Resource Ask" not in text
    assert "Not enough history" in text


def _pdf_text(pdf_bytes: bytes) -> str:
    """Extract visible text from a fpdf-generated PDF (flate streams)."""
    import re
    import zlib

    texts = []
    for m in re.finditer(rb"stream\r?\n(.*?)endstream", pdf_bytes, re.S):
        try:
            texts.append(zlib.decompress(m.group(1)).decode("latin-1"))
        except Exception:
            pass
    return "\n".join(texts)


# ── OSCAL POA&M export ────────────────────────────────────────────────────────


def test_oscal_export_structure():
    import json
    import uuid

    from poam_export import export_poam_oscal

    answers = all_met()
    answers["AC.L2-3.1.1"] = {
        "status": "NOT MET",
        "owner": "jane",
        "target_date": "2026-12-31",
        "remediation_plan": "Implement access controls",
        "likelihood": "High",
        "impact": "Medium",
        "estimated_cost": "5000",
        "assessor_notes": "No MFA on admin accounts",
    }
    doc = export_poam_oscal(answers, SCOPED, org_name="Test Org")

    poam = doc["plan-of-action-and-milestones"]
    assert uuid.UUID(poam["uuid"])
    assert poam["metadata"]["oscal-version"] == "1.1.2"
    assert poam["metadata"]["title"] == "POA&M — Test Org"
    assert poam["poam-items"]

    item = poam["poam-items"][0]
    assert uuid.UUID(item["uuid"])
    props = {p["name"]: p["value"] for p in item["props"]}
    assert props["control-id"] == "AC.L2-3.1.1"
    assert props["poam-id"].startswith("WK-")
    assert props["poc"] == "jane"
    assert props["target-date"] == "2026-12-31"
    assert item["remediation"]["purpose"] == "Implement access controls"


def test_oscal_export_empty():
    from poam_export import export_poam_oscal

    doc = export_poam_oscal({}, [], org_name="Empty Org")
    poam = doc["plan-of-action-and-milestones"]
    assert poam["poam-items"] == []
    assert poam["metadata"]["title"] == "POA&M — Empty Org"
