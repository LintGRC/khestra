#!/usr/bin/env python3
"""Generate the four-file SSP Appendix Starter Pack."""

from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import Path

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "appendix_pack"

README = """SSP Appendix Starter Pack
=========================

These are structural templates — not approved policies or audit evidence until your
organization completes, reviews, and signs them.

FILES
-----
1. IR_Plan_Outline.docx          Incident response plan framework (phases, contacts, reporting)
2. Access_Control_Matrix.xlsx    Roles + mapping to Access Control (AC) requirement IDs
3. Asset_Inventory.xlsx          Hardware and software inventory aligned to assessment boundary
4. Network_Topology.drawio       Editable diagram shell (open in diagrams.net / draw.io)

HOW TO USE WITH THE CMMC TOOL
-----------------------------
- Complete inventories → copy summary into System Profile or attach full spreadsheet
- Finish topology in draw.io → export PNG → upload under Organization → Topology
- Reference completed documents in SSP narratives and appendix index
- Have management review before C3PAO or self-assessment

DISCLAIMER
----------
Drafting utility only. Does not provide legal advice or guarantee compliance.
"""


def _ir_plan_outline(path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading("Incident Response Plan — Outline", level=0)
    doc.add_paragraph("Organization: [COMPANY NAME]")
    doc.add_paragraph("System: [SYSTEM NAME]")
    doc.add_paragraph("Version: [0.1 DRAFT]    Effective date: [DATE]    Next review: [DATE]")
    doc.add_paragraph(
        "Status: DRAFT — requires review and approval by [SYSTEM OWNER / ISO] before operational use."
    )

    sections = [
        (
            "1. Purpose and scope",
            [
                "Describe systems and CUI in scope for this plan.",
                "Reference the System Security Plan section on incident response.",
            ],
        ),
        (
            "2. Roles and contact roster",
            [
                "Incident Commander: [NAME, TITLE, PHONE, EMAIL]",
                "ISO / Security Lead: [NAME]",
                "IT Operations: [NAME]",
                "Legal / Management notification: [NAME]",
                "External reporting (DoW / DIB): [PROCESS OWNER]",
            ],
        ),
        (
            "3. Incident phases",
            [
                "Detection & reporting — how users report suspected incidents.",
                "Analysis & classification — CUI spillage, malware, unauthorized access, etc.",
                "Containment — isolate systems, preserve evidence.",
                "Eradication & recovery — restore from backup, validate integrity.",
                "Post-incident — lessons learned, POA&M updates, management briefing.",
            ],
        ),
        (
            "4. Reporting and notification timelines",
            [
                "Internal escalation within [X] hours.",
                "DoW / contract requirements per [CONTRACT CLAUSE / DFARS reference].",
                "Law enforcement / cyber center contacts as applicable.",
            ],
        ),
        (
            "5. Testing and maintenance",
            [
                "Tabletop exercise frequency: [e.g. annual].",
                "Plan review cycle: [e.g. annual or after major change].",
            ],
        ),
        (
            "6. Approval",
            [
                "Prepared by: _________________________  Date: _________",
                "Approved by (System Owner): __________  Date: _________",
            ],
        ),
    ]
    for title, bullets in sections:
        doc.add_heading(title, level=1)
        for line in bullets:
            doc.add_paragraph(line, style="List Bullet")

    doc.save(path)


def _access_control_matrix(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from controls import CMMC_FRAMEWORK

    wb = Workbook()
    ws_roles = wb.active
    ws_roles.title = "Roles"
    role_headers = [
        "Role title",
        "Person name",
        "Email",
        "Privileged (Y/N)",
        "Systems / applications",
        "Access level summary",
        "Last access review",
        "Reviewer",
    ]
    ws_roles.append(role_headers)
    for h in role_headers:
        ws_roles.cell(1, role_headers.index(h) + 1).font = Font(bold=True)
    sample_roles = [
        ("System Owner", "[NAME]", "[email]", "Y", "[ERP, M365 admin]", "Full admin", "[DATE]", "[ISO]"),
        ("Standard user", "[NAME]", "[email]", "N", "[SharePoint CUI]", "Read/write CUI library", "[DATE]", "[IT]"),
        ("Helpdesk", "[NAME]", "[email]", "Y", "[Entra ID, Intune]", "User support, no CUI export", "[DATE]", "[ISO]"),
    ]
    for row in sample_roles:
        ws_roles.append(list(row))

    ws_map = wb.create_sheet("AC_Control_Map")
    map_headers = [
        "Control ID",
        "SPRS weight",
        "Requirement (abbreviated)",
        "Responsible role",
        "Implementation reference",
        "Evidence / artifact",
    ]
    ws_map.append(map_headers)
    for c, h in enumerate(map_headers, 1):
        ws_map.cell(1, c).font = Font(bold=True)
    for cid, info in sorted(CMMC_FRAMEWORK.items()):
        if info["family"] != "Access Control":
            continue
        ws_map.append(
            [
                cid,
                info["weight"],
                info["name"][:120],
                "[ROLE]",
                "[Policy / config reference]",
                "[Ticket, export, or log name]",
            ]
        )

    for sheet in (ws_roles, ws_map):
        for col in sheet.columns:
            letter = col[0].column_letter
            width = min(max(max(len(str(cell.value or "")) for cell in col) + 2, 12), 50)
            sheet.column_dimensions[letter].width = width

    wb.save(path)


def _asset_inventory(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    hw = wb.active
    hw.title = "Hardware"
    hw_headers = [
        "Asset ID",
        "Asset name",
        "Type",
        "Location",
        "In CMMC boundary (Y/N)",
        "Processes CUI (Y/N)",
        "Owner",
        "OS / firmware",
        "Last patch date",
        "Notes",
    ]
    hw.append(hw_headers)
    for c, h in enumerate(hw_headers, 1):
        hw.cell(1, c).font = Font(bold=True)
    hw.append(
        [
            "HW-001",
            "[Laptop model]",
            "Endpoint",
            "[Office / remote]",
            "Y",
            "Y",
            "[USER]",
            "[Windows 11]",
            "[DATE]",
            "",
        ]
    )

    sw = wb.create_sheet("Software")
    sw_headers = [
        "Asset ID",
        "Software name",
        "Version",
        "Hosting",
        "In boundary (Y/N)",
        "CUI relevance",
        "License owner",
        "Notes",
    ]
    sw.append(sw_headers)
    for c, h in enumerate(sw_headers, 1):
        sw.cell(1, c).font = Font(bold=True)
    sw.append(
        [
            "SW-001",
            "[Microsoft 365 GCC]",
            "[version]",
            "Cloud",
            "Y",
            "CUI storage / collaboration",
            "[IT]",
            "",
        ]
    )

    for sheet in (hw, sw):
        for col in sheet.columns:
            letter = col[0].column_letter
            width = min(max(max(len(str(cell.value or "")) for cell in col) + 2, 10), 40)
            sheet.column_dimensions[letter].width = width

    wb.save(path)


def _network_topology_drawio(path: Path) -> None:
    """Minimal draw.io diagram — open at https://app.diagrams.net"""
    xml = """<?xml version="1.0" encoding="UTF-8"?>
<mxfile host="app.diagrams.net" modified="2026-01-01T00:00:00.000Z" agent="CMMC Appendix Pack" version="22.1.0">
  <diagram id="topology" name="Network Topology">
    <mxGraphModel dx="1200" dy="800" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="1100" pageHeight="850" math="0" shadow="0">
      <root>
        <mxCell id="0"/>
        <mxCell id="1" parent="0"/>
        <mxCell id="title" value="[COMPANY NAME] — CUI System Boundary (edit in draw.io)" style="text;html=1;fontSize=16;fontStyle=1;align=left;" vertex="1" parent="1">
          <mxGeometry x="40" y="20" width="600" height="30" as="geometry"/>
        </mxCell>
        <mxCell id="users" value="Users&#xa;[count / roles]" style="rounded=1;whiteSpace=wrap;html=1;fillColor=#dae8fc;strokeColor=#6c8ebf;" vertex="1" parent="1">
          <mxGeometry x="40" y="100" width="140" height="70" as="geometry"/>
        </mxCell>
        <mxCell id="idp" value="Identity / MFA&#xa;[Entra ID / AD]" style="rounded=1;whiteSpace=wrap;html=1;fillColor=#dae8fc;strokeColor=#6c8ebf;" vertex="1" parent="1">
          <mxGeometry x="240" y="90" width="160" height="80" as="geometry"/>
        </mxCell>
        <mxCell id="cui" value="CUI enclave&#xa;[SharePoint / app]" style="rounded=1;whiteSpace=wrap;html=1;fillColor=#fff2cc;strokeColor=#d6b656;" vertex="1" parent="1">
          <mxGeometry x="460" y="90" width="160" height="80" as="geometry"/>
        </mxCell>
        <mxCell id="endpoint" value="Endpoint mgmt&#xa;[Intune / AV]" style="rounded=1;whiteSpace=wrap;html=1;fillColor=#d5e8d4;strokeColor=#82b366;" vertex="1" parent="1">
          <mxGeometry x="680" y="100" width="150" height="70" as="geometry"/>
        </mxCell>
        <mxCell id="fw" value="Firewall / VPN&#xa;[vendor]" style="rounded=1;whiteSpace=wrap;html=1;fillColor=#f8cecc;strokeColor=#b85450;" vertex="1" parent="1">
          <mxGeometry x="240" y="240" width="160" height="70" as="geometry"/>
        </mxCell>
        <mxCell id="siem" value="Logging / SIEM&#xa;[retention period]" style="rounded=1;whiteSpace=wrap;html=1;fillColor=#e1d5e7;strokeColor=#9673a6;" vertex="1" parent="1">
          <mxGeometry x="460" y="240" width="160" height="70" as="geometry"/>
        </mxCell>
        <mxCell id="oos" value="OUT OF SCOPE&#xa;[HR / payroll / etc.]" style="rounded=1;whiteSpace=wrap;html=1;dashed=1;fillColor=#f5f5f5;strokeColor=#666666;" vertex="1" parent="1">
          <mxGeometry x="40" y="360" width="200" height="60" as="geometry"/>
        </mxCell>
        <mxCell id="note" value="Export as PNG when complete → upload in CMMC tool → Organization → Topology" style="text;html=1;fontSize=11;align=left;fontColor=#666666;" vertex="1" parent="1">
          <mxGeometry x="40" y="450" width="500" height="40" as="geometry"/>
        </mxCell>
        <mxCell id="e1" style="endArrow=classic;html=1;" edge="1" parent="1" source="users" target="idp">
          <mxGeometry relative="1" as="geometry"/>
        </mxCell>
        <mxCell id="e2" style="endArrow=classic;html=1;" edge="1" parent="1" source="idp" target="cui">
          <mxGeometry relative="1" as="geometry"/>
        </mxCell>
        <mxCell id="e3" style="endArrow=classic;html=1;" edge="1" parent="1" source="cui" target="endpoint">
          <mxGeometry relative="1" as="geometry"/>
        </mxCell>
        <mxCell id="e4" style="endArrow=classic;html=1;" edge="1" parent="1" source="idp" target="fw">
          <mxGeometry relative="1" as="geometry"/>
        </mxCell>
        <mxCell id="e5" style="endArrow=classic;html=1;" edge="1" parent="1" source="fw" target="siem">
          <mxGeometry relative="1" as="geometry"/>
        </mxCell>
      </root>
    </mxGraphModel>
  </diagram>
</mxfile>
"""
    path.write_text(xml.strip() + "\n", encoding="utf-8")


def generate_all(out_dir: Path | None = None) -> Path:
    out = out_dir or OUT_DIR
    out.mkdir(parents=True, exist_ok=True)
    (out / "README.txt").write_text(README, encoding="utf-8")
    _ir_plan_outline(out / "IR_Plan_Outline.docx")
    _access_control_matrix(out / "Access_Control_Matrix.xlsx")
    _asset_inventory(out / "Asset_Inventory.xlsx")
    _network_topology_drawio(out / "Network_Topology.drawio")
    return out


def build_appendix_pack_zip(out_dir: Path | None = None) -> bytes:
    folder = generate_all(out_dir)
    buf = BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in sorted(folder.iterdir()):
            if path.is_file():
                zf.write(path, arcname=f"appendix_pack/{path.name}")
    buf.seek(0)
    return buf.getvalue()


