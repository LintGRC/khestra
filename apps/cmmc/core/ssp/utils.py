"""Shared helpers for SSP section builders."""

from docx.shared import RGBColor
from controls import CMMC_FRAMEWORK


def sanitize_answers(answers: dict) -> dict:
    """Normalize assessment records for document generation."""
    sanitized = {}
    for cid, data in answers.items():
        if data is None:
            sanitized[cid] = _empty_answer()
            continue
        narrative = (
            data.get("implementation_narrative", "").strip()
            or data.get("implementation_desc", "").strip()
            or data.get("assessor_notes", "").strip()
            or data.get("remediation_plan", "").strip()
        )
        sanitized[cid] = {
            "status": data.get("status", "NOT STARTED"),
            "implementation_narrative": narrative,
            "implementation_desc": narrative,
            "examine": data.get("examine", ""),
            "interview": data.get("interview", ""),
            "test": data.get("test", ""),
            "assessor_notes": data.get("assessor_notes", ""),
            "owner": data.get("owner", ""),
            "target_date": data.get("target_date", ""),
            "estimated_cost": data.get("estimated_cost", ""),
            "evidence": data.get("evidence", []),
            "linked_policies": data.get("linked_policies", []),
            "linked_assets": data.get("linked_assets", []),
            "linked_team": data.get("linked_team", []),
        }
    return sanitized


def _empty_answer() -> dict:
    return {
        "status": "NOT STARTED",
        "implementation_narrative": "",
        "implementation_desc": "",
        "examine": "",
        "interview": "",
        "test": "",
        "assessor_notes": "",
        "owner": "",
        "target_date": "",
        "estimated_cost": "",
        "evidence": [],
    }


def group_controls_by_family(answers: dict, scoped_controls: list) -> dict:
    """Return {family: [(control_id, answer), ...]} for scoped controls only."""
    sanitized = sanitize_answers(answers)
    families: dict = {}
    for cid in scoped_controls:
        if cid not in CMMC_FRAMEWORK or cid not in sanitized:
            continue
        family = CMMC_FRAMEWORK[cid]["family"]
        families.setdefault(family, []).append((cid, sanitized[cid]))
    for family in families:
        families[family].sort(key=lambda item: item[0])
    return families


def add_placeholder(doc, text: str) -> None:
    para = doc.add_paragraph()
    run = para.add_run(f"[PLACEHOLDER: {text}]")
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.bold = True
    run.italic = True


def create_table(doc, headers: list, data_rows: list):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = str(header)
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
    for row_data in data_rows:
        row_cells = table.add_row().cells
        for i, cell_data in enumerate(row_data):
            row_cells[i].text = str(cell_data)
    return table


def control_ssp_supplement(ans: dict) -> str:
    """Evidence filename line appended to CUI SSP narrative cells (audit package crosswalk)."""
    evidence_list = ans.get("evidence") or []
    if not evidence_list:
        return ""
    names = "; ".join(ev.get("filename", "unknown") for ev in evidence_list)
    return f"Attached evidence (audit package): {names}"


def risk_severity(weight: int) -> str:
    if weight >= 5:
        return "Critical"
    if weight == 3:
        return "High"
    if weight == 1:
        return "Moderate"
    return "Low"
