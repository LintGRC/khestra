"""Assemble SSP sections into a single Word document."""

import os
from io import BytesIO

from docx import Document

from sprs_engine import calculate_detailed_sprs
from ssp.constants import FAMILY_ORDER
from ssp.doc_styles import patch_document
from ssp.export_notice import append_workspace_snapshot_notice
from ssp.template_utils import clear_document_body, is_official_cui_template
from ssp.template_filler import fill_cui_ssp_template
from ssp.sections.asset_inventory_appendix import add_asset_inventory_appendix
from ssp.sections.record_of_changes import add_record_of_changes
from ssp.sections.control_family import add_control_family_section, add_control_narratives_intro
from ssp.sections.cover import add_cover_page
from ssp.sections.evidence import add_evidence_appendix
from ssp.sections.policy_appendix import add_policy_appendix
from ssp.sections.executive_summary import add_executive_summary
from ssp.sections.poam import add_poam_section
from ssp.sections.roles import add_roles_responsibilities
from ssp.sections.system_identification import add_system_identification
from ssp.sections.system_overview import add_system_overview
from ssp.sections.toc import add_table_of_contents
from ssp.utils import group_controls_by_family, sanitize_answers


class SSPAssembler:
    """Build SSP one section at a time; control families are separate modular blocks."""

    def __init__(self, org_name: str, template_path: str | None = None):
        self.template_path = template_path
        self.uses_cui_template = is_official_cui_template(template_path)
        if template_path and os.path.exists(template_path):
            self.doc = Document(template_path)
        else:
            self.doc = Document()
        if not self.uses_cui_template:
            patch_document(self.doc)
        self.org_name = org_name

    def add_front_matter(
        self,
        answers: dict,
        asset_scope: dict,
        scoped_controls: list,
        org_profile: dict | None = None,
        sprs_result: dict | None = None,
        system_scope: dict | None = None,
        ssp_version: str = "",
        prepared_by: str = "",
        approved_by: str = "",
        org_asset_bytes: dict | None = None,
    ) -> None:
        profile = org_profile or {"org_name": self.org_name}
        org_label = profile.get("org_name") or self.org_name
        add_cover_page(self.doc, org_label, ssp_version=ssp_version, prepared_by=prepared_by, approved_by=approved_by)
        add_table_of_contents(self.doc, cui_outline=self.uses_cui_template)
        if self.uses_cui_template:
            add_system_identification(
                self.doc,
                org_profile=profile,
                asset_scope=asset_scope,
                answers=answers,
                scoped_controls=scoped_controls,
                sprs_result=sprs_result,
            )
        else:
            add_executive_summary(
                self.doc,
                sanitize_answers(answers),
                asset_scope,
                scoped_controls,
                org_profile=profile,
                sprs_result=sprs_result,
            )
        add_system_overview(self.doc, asset_scope, org_profile=profile, cui_outline=self.uses_cui_template, system_scope=system_scope, org_asset_bytes=org_asset_bytes)
        add_roles_responsibilities(self.doc, org_profile=profile)

    def add_all_control_families(
        self,
        answers: dict,
        scoped_controls: list,
        risks_by_control: dict | None = None,
        org_profile: dict | None = None,
        org_inventory: dict | None = None,
    ) -> None:
        section_number = 5
        add_control_narratives_intro(
            self.doc, section_number=section_number, cui_outline=self.uses_cui_template
        )
        grouped = group_controls_by_family(answers, scoped_controls)
        active_families = [f for f in FAMILY_ORDER if f in grouped]
        from poam_export import poam_weakness_ids
        poam_ids = poam_weakness_ids(answers, scoped_controls)
        for idx, family in enumerate(active_families, 1):
            add_control_family_section(
                self.doc,
                idx,
                family,
                grouped[family],
                section_number=section_number,
                page_break_after=idx < len(active_families),
                risks_by_control=risks_by_control,
                org_profile=org_profile,
                org_inventory=org_inventory,
                poam_ids=poam_ids,
            )

    def add_poam_before_controls(self, answers: dict, scoped_controls: list) -> None:
        add_poam_section(self.doc, answers, scoped_controls, section_number=4)

    def add_back_matter(
        self,
        answers: dict,
        scoped_controls: list,
        audit_log: list | None = None,
        org_inventory: dict | None = None,
    ) -> None:
        add_evidence_appendix(self.doc, answers, scoped_controls, section_number=6)
        add_policy_appendix(self.doc, scoped_controls, section_number=7)
        add_asset_inventory_appendix(self.doc, org_inventory)
        add_record_of_changes(self.doc, audit_log)
        append_workspace_snapshot_notice(self.doc)

    def assemble(
        self,
        answers: dict,
        asset_scope: dict,
        scoped_controls: list | None = None,
        org_profile: dict | None = None,
        org_assets: dict | None = None,
        org_asset_bytes: dict | None = None,
        audit_log: list | None = None,
        org_inventory: dict | None = None,
        risks_by_control: dict | None = None,
        system_scope: dict | None = None,
        ssp_version: str = "",
        prepared_by: str = "",
        approved_by: str = "",
    ) -> bytes:
        scoped = scoped_controls or list(answers.keys())
        profile = org_profile or {"org_name": self.org_name}

        if self.uses_cui_template:
            fill_cui_ssp_template(
                self.doc,
                answers,
                asset_scope,
                scoped,
                profile,
                org_assets=org_assets,
                org_asset_bytes=org_asset_bytes,
                audit_log=audit_log,
                org_inventory=org_inventory,
                risks_by_control=risks_by_control,
                system_scope=system_scope,
                ssp_version=ssp_version,
                prepared_by=prepared_by,
                approved_by=approved_by,
            )
        else:
            sprs_result = calculate_detailed_sprs(answers, scoped)
            self.add_front_matter(
                answers, asset_scope, scoped, org_profile=profile, sprs_result=sprs_result,
                system_scope=system_scope, ssp_version=ssp_version,
                prepared_by=prepared_by, approved_by=approved_by,
                org_asset_bytes=org_asset_bytes,
            )
            self.add_poam_before_controls(answers, scoped)
            self.add_all_control_families(
                answers, scoped, risks_by_control=risks_by_control,
                org_profile=profile, org_inventory=org_inventory,
            )
            self.add_back_matter(answers, scoped, audit_log=audit_log, org_inventory=org_inventory)

        buffer = BytesIO()
        self.doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


def assemble_ssp(
    answers: dict,
    asset_scope: dict,
    org_name: str = "Your Organization",
    scoped_controls: list | None = None,
    template_path: str | None = None,
    org_profile: dict | None = None,
    org_assets: dict | None = None,
    org_asset_bytes: dict | None = None,
    audit_log: list | None = None,
    org_inventory: dict | None = None,
    risks_by_control: dict | None = None,
    system_scope: dict | None = None,
    ssp_version: str = "",
    prepared_by: str = "",
    approved_by: str = "",
) -> bytes:
    """Public API: build full SSP from modular sections."""
    profile = org_profile or {"org_name": org_name}
    if not profile.get("org_name"):
        profile["org_name"] = org_name
    assembler = SSPAssembler(profile.get("org_name") or org_name, template_path)
    return assembler.assemble(
        answers,
        asset_scope,
        scoped_controls,
        org_profile=profile,
        org_assets=org_assets,
        org_asset_bytes=org_asset_bytes,
        audit_log=audit_log,
        org_inventory=org_inventory,
        risks_by_control=risks_by_control,
        system_scope=system_scope,
        ssp_version=ssp_version,
        prepared_by=prepared_by,
        approved_by=approved_by,
    )


def assemble_ssp_family(
    answers: dict,
    family_name: str,
    org_name: str = "Your Organization",
    scoped_controls: list | None = None,
    risks_by_control: dict | None = None,
    org_profile: dict | None = None,
    org_inventory: dict | None = None,
) -> bytes:
    """Export a single control family as a standalone document (for iterative editing)."""
    scoped = scoped_controls or list(answers.keys())
    grouped = group_controls_by_family(answers, scoped)
    if family_name not in grouped:
        raise ValueError(f"No scoped controls in family: {family_name}")

    doc = Document()
    patch_document(doc)
    doc.add_heading(f"SSP Section — {family_name}", level=0)
    doc.add_paragraph(f"Organization: {org_name}")
    family_index = FAMILY_ORDER.index(family_name) + 1 if family_name in FAMILY_ORDER else 1
    from poam_export import poam_weakness_ids
    add_control_family_section(
        doc, family_index, family_name, grouped[family_name],
        page_break_after=False, risks_by_control=risks_by_control,
        org_profile=org_profile, org_inventory=org_inventory,
        poam_ids=poam_weakness_ids(answers, scoped),
    )

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
