"""Safe heading and list styles for the official DoW CUI SSP Word template.

The public CUI SSP .docx often defines custom styles (Heading 2, WP9_Header, etc.)
instead of built-in Word names (Heading 1, List Bullet). These helpers avoid
export failures while preserving template typography where possible.
"""

from docx.shared import Pt

HEADING_STYLE_FALLBACKS = {
    0: ("Title", "Heading 2", "Normal"),
    1: ("Heading 1", "Heading 2", "Heading 4"),
    2: ("Heading 2", "Heading 4", "Normal"),
    3: ("Heading 3", "Heading 4", "Normal"),
    4: ("Heading 4", "Normal"),
}

BULLET_STYLE_FALLBACKS = ("List Bullet", "List Paragraph", "Normal")


def _style_exists(doc, name: str) -> bool:
    try:
        doc.styles[name]
        return True
    except KeyError:
        return False


def safe_add_heading(doc, text: str, level: int = 1):
    for style in HEADING_STYLE_FALLBACKS.get(level, ("Normal",)):
        if _style_exists(doc, style):
            return doc.add_paragraph(text, style=style)
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = True
    run.font.size = Pt(16 if level <= 1 else 14)
    return para


def safe_add_bullet(doc, text: str, add_paragraph=None):
    if add_paragraph is None:
        add_paragraph = doc.add_paragraph
    for style in BULLET_STYLE_FALLBACKS:
        if _style_exists(doc, style):
            return add_paragraph(text, style=style)
    return add_paragraph(f"• {text}")


def patch_document(doc) -> None:
    """Replace add_heading with template-safe version; sections use add_heading as usual."""

    def patched_add_heading(text: str, level: int = 1):
        return safe_add_heading(doc, text, level=level)

    doc.add_heading = patched_add_heading  # type: ignore[method-assign]

    orig_add_paragraph = doc.add_paragraph

    def patched_add_paragraph(text: str = "", style=None):
        if style == "List Bullet":
            return safe_add_bullet(doc, text, add_paragraph=orig_add_paragraph)
        return orig_add_paragraph(text, style=style)

    doc.add_paragraph = patched_add_paragraph  # type: ignore[method-assign]
