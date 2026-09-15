from datetime import datetime

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from ssp.export_notice import add_cover_snapshot_notice


def add_cover_page(doc, org_name: str, ssp_version: str = "", prepared_by: str = "", approved_by: str = "") -> None:
    doc.add_paragraph()
    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = notice.add_run("NOTICE: This document may contain Controlled Unclassified Information (CUI)")
    run.bold = True
    run.font.size = Pt(10)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Handle in accordance with NIST SP 800-171 requirements").font.size = Pt(9)

    doc.add_paragraph()
    title = doc.add_heading("SYSTEM SECURITY PLAN", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    org_para = doc.add_paragraph()
    org_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    org_para.add_run(org_name).font.size = Pt(18)

    doc.add_paragraph()
    sys_para = doc.add_paragraph()
    sys_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sys_para.add_run("CMMC Level 2 Assessment System").font.size = Pt(16)

    doc.add_paragraph()
    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_para.add_run(f"Prepared: {datetime.now().strftime('%B %d, %Y')}").font.size = Pt(14)

    if ssp_version:
        ver_para = doc.add_paragraph()
        ver_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ver_para.add_run(f"Version {ssp_version}").font.size = Pt(12)

    if prepared_by:
        pre_para = doc.add_paragraph()
        pre_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pre_para.add_run(f"Prepared by: {prepared_by}").font.size = Pt(11)

    if approved_by:
        app_para = doc.add_paragraph()
        app_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        app_para.add_run(f"Approved by: {approved_by}").font.size = Pt(11)

    add_cover_snapshot_notice(doc)

    doc.add_page_break()
