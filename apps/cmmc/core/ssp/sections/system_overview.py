from ssp.utils import add_placeholder, create_table


def add_system_overview(doc, asset_scope: dict, org_profile: dict | None = None, *, cui_outline: bool = False, system_scope: dict | None = None, org_asset_bytes: dict | None = None) -> None:
    profile = org_profile or {}
    scope = system_scope or {}
    section_title = "2. System Environment" if cui_outline else "2. System Overview"
    doc.add_heading(section_title, level=1)

    if profile.get("system_description", "").strip():
        doc.add_paragraph(profile["system_description"].strip())
    else:
        doc.add_paragraph(
            "This section describes the organizational system(s) subject to CMMC assessment."
        )

    doc.add_heading("2.1 System Architecture", level=2)
    if profile.get("architecture_summary", "").strip():
        doc.add_paragraph(profile["architecture_summary"].strip())
    else:
        add_placeholder(
            doc,
            "Describe system components, network topology, data flows, and security control implementation points.",
        )

    doc.add_heading("2.2 Data Classification", level=2)
    scope_data = [
        ["CUI Assets", "Controls fully apply", f"{asset_scope.get('CUI Assets', 0)}%"],
        ["Security Protection Assets", "Controls apply to protection mechanisms", f"{asset_scope.get('Security Protection Assets', 0)}%"],
        ["Contractor Risk Managed Assets", "Limited control application", f"{asset_scope.get('Contractor Risk Managed Assets', 0)}%"],
        ["Out-of-Scope Assets", "Controls do not apply", f"{asset_scope.get('Out-of-Scope Assets', 0)}%"],
    ]
    create_table(doc, ["Asset Type", "Description", "Scope Percentage"], scope_data)

    if scope.get("cui_types", "").strip():
        doc.add_heading("2.2.1 CUI Types", level=3)
        doc.add_paragraph(scope["cui_types"].strip())

    doc.add_heading("2.3 System Boundaries", level=2)
    if profile.get("boundary_description", "").strip():
        doc.add_paragraph(profile["boundary_description"].strip())
    else:
        add_placeholder(
            doc,
            "Define logical and physical boundaries, network segments, cloud environments, and third-party integrations.",
        )

    _embed_topology(doc, org_asset_bytes)

    doc.add_heading("2.4 External Connections", level=2)
    conns = scope.get("external_connections", "").strip()
    if conns:
        for line in conns.splitlines():
            if line.strip():
                doc.add_paragraph(line.strip(), style="List Bullet")
    else:
        add_placeholder(
            doc,
            "List external system interconnections, APIs, VPNs, and third-party integrations authorized to access the system.",
        )

    if scope.get("last_review", "").strip():
        doc.add_heading("2.5 Scope Review", level=2)
        doc.add_paragraph(f"System scope last reviewed: {scope['last_review'].strip()}")

    doc.add_page_break()


def _embed_topology(doc, org_asset_bytes: dict | None) -> None:
    """Embed the uploaded topology diagram image in the document."""
    if not org_asset_bytes:
        return
    try:
        from org_assets import get_topology_bytes
    except ImportError:
        return

    topo = get_topology_bytes({"topology_filename": ""}, org_asset_bytes)
    if not topo:
        return

    from io import BytesIO
    from docx.shared import Inches

    try:
        image_stream = BytesIO(topo[1])
        caption = doc.add_paragraph()
        caption.add_run("Figure 1: CUI Enclave Network Topology").italic = True
        doc.add_picture(image_stream, width=Inches(5.5))
        doc.add_paragraph()
    except Exception:
        pass
