"""Table of contents — must match generated section order."""

from docx.enum.text import WD_ALIGN_PARAGRAPH

from ssp.constants import FAMILY_ORDER


def add_table_of_contents(doc, *, cui_outline: bool = False) -> None:
    heading = doc.add_heading("Table of Contents", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if cui_outline:
        items = [
            "1. System Identification",
            "2. System Environment",
            "3. Roles and Responsibilities",
            "4. Plan of Action and Milestones (POA&M)",
            "5. Security Requirements (NIST SP 800-171 Rev 2)",
        ]
        for idx, family in enumerate(FAMILY_ORDER, 1):
            items.append(f"   5.{idx} {family}")
        items.append("6. Evidence Appendix")
    else:
        items = [
            "1. Executive Summary",
            "2. System Overview",
            "3. Roles and Responsibilities",
            "4. Plan of Action and Milestones (POA&M)",
            "5. Control Implementation Narratives",
        ]
        for idx, family in enumerate(FAMILY_ORDER, 1):
            items.append(f"   5.{idx} {family}")
        items.append("6. Evidence Appendix")

    for item in items:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()
