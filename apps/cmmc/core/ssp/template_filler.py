"""Fill the official DoW CUI SSP template in place — preserves layout and styles."""

from __future__ import annotations

import re
from datetime import datetime
from difflib import SequenceMatcher
from io import BytesIO
from typing import Any, Dict, List, Optional

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from controls import CMMC_FRAMEWORK
from org_assets import get_topology_bytes
from ssp.constants import FAMILY_ORDER
from org_inventory import inventory_summary
from ssp.export_notice import append_workspace_snapshot_notice
from ssp.utils import control_ssp_supplement
from ssp.sections.asset_inventory_appendix import add_asset_inventory_appendix
from ssp.sections.record_of_changes import add_record_of_changes

_PLACEHOLDER_DETAIL = (
    "Current implementation or planned implementation details.  "
    "Include a detailed description of how the control is implemented, "
    "or the plan to implement the control."
)
_FAMILY_HEADERS = set(FAMILY_ORDER) | {
    "Access Control",
    "Awareness and Training",
    "Audit and Accountability",
    "Configuration Management",
    "Identification and Authentication",
    "Incident Response",
    "Maintenance",
    "Media Protection",
    "Personnel Security",
    "Physical Protection",
    "Risk Assessment",
    "Security Assessment",
    "System and Communications Protection",
    "System and Information Integrity",
}


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (text or "").lower()).strip()


def _match_control_id(requirement_text: str) -> Optional[str]:
    req = _norm(requirement_text)
    req_tokens = set(req.split())
    best_cid = None
    best_score = 0.0
    for cid, info in CMMC_FRAMEWORK.items():
        name = _norm(info["name"])
        name_tokens = set(name.split())
        overlap = len(req_tokens & name_tokens) / max(len(req_tokens), 1)
        seq = SequenceMatcher(None, req[:90], name[:90]).ratio()
        score = max(overlap * 0.65 + seq * 0.35, seq)
        if req[:35] and req[:35] in name:
            score = max(score, 0.88)
        if score > best_score:
            best_score = score
            best_cid = cid
    return best_cid if best_score >= 0.42 else None


def _set_cell(table, row: int, col: int, text: str) -> None:
    if row >= len(table.rows):
        return
    cells = table.rows[row].cells
    if col >= len(cells):
        return
    cells[col].text = text


def _split_name_title(value: str) -> tuple[str, str]:
    value = (value or "").strip()
    if "," in value:
        name, title = value.split(",", 1)
        return name.strip(), title.strip()
    return value, ""


def _run_has_drawing(run) -> bool:
    el = run._element
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return el.find(f".//{w}drawing") is not None or el.find(f".//{w}pict") is not None


def _resolve_header_name(org_profile: dict | None, org_name: str) -> str:
    """Pick a header label that fits — full org name in body, shorter line in page header."""
    profile = org_profile or {}
    short = (profile.get("header_short_name") or "").strip()
    if short:
        return short
    org = (org_name or profile.get("org_name") or "Organization").strip()
    system = (profile.get("system_name") or "").strip()
    if len(org) > 32 and system and len(system) <= 42:
        return system
    return org


def _apply_header_style(para, doc) -> None:
    try:
        para.style = doc.styles["Header"]
    except KeyError:
        pass


def _set_title_date_line(para, date_str: str) -> None:
    """Second header line: SSP title left, last-updated right (tab stop)."""
    from docx.enum.text import WD_TAB_ALIGNMENT, WD_PARAGRAPH_ALIGNMENT
    from docx.shared import Inches

    para.text = ""
    para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    para.paragraph_format.tab_stops.clear_all()
    para.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    r1 = para.add_run("SYSTEM SECURITY PLAN")
    r1.bold = True
    r2 = para.add_run(f"\tLast Updated: {date_str}")
    r2.bold = True


def fix_page_headers(
    doc,
    org_name: str,
    updated_date: str | None = None,
    org_profile: dict | None = None,
) -> None:
    """Fill header placeholders; use two lines when the organization name is long."""
    name = _resolve_header_name(org_profile, org_name)
    date_str = updated_date or datetime.now().strftime("%B %d, %Y")
    use_two_lines = len(name) > 24

    for section in doc.sections:
        header = section.header
        if header is None or not header.paragraphs:
            continue

        para = header.paragraphs[0]
        _apply_header_style(para, doc)
        text_runs = [r for r in para.runs if not _run_has_drawing(r)]

        if use_two_lines:
            while len(header.paragraphs) > 1:
                el = header.paragraphs[-1]._element
                el.getparent().remove(el)
            if text_runs:
                text_runs[0].text = name
                for r in text_runs[1:]:
                    r.text = ""
                for r in text_runs:
                    r.bold = True
                    r.font.size = None
            p2 = header.add_paragraph()
            _apply_header_style(p2, doc)
            _set_title_date_line(p2, date_str)
            continue

        while len(header.paragraphs) > 1:
            el = header.paragraphs[-1]._element
            el.getparent().remove(el)

        if len(text_runs) >= 5:
            text_runs[0].text = name
            if len(text_runs) > 1:
                text_runs[1].text = ""
            if len(text_runs) > 2:
                text_runs[2].text = ""
            text_runs[3].text = " SYSTEM SECURITY PLAN   "
            text_runs[4].text = f"\t                    Last Updated:  {date_str}"
            for r in text_runs[5:]:
                r.text = ""
            for r in text_runs:
                r.bold = True
                r.font.size = None
        else:
            for r in text_runs:
                t = r.text
                if "Insert name" in t or t.strip() in ("<<", ">>", "name", "Insert "):
                    r.text = name if "name" in t.lower() or t == "<<" else ""
                elif "Insert date" in t:
                    r.text = date_str
                elif "Last Updated" in t:
                    r.text = f"\t                    Last Updated:  {date_str}"
                elif t.strip() in (">>", "<<"):
                    r.text = ""
                r.bold = True
                r.font.size = None


def _replace_bracket(text: str, narrative: str) -> str:
    if not narrative:
        return text
    return re.sub(r"\[[^\]]+\]", narrative, text, count=1)


def _fill_identification_paragraphs(doc, org_profile: dict, asset_scope: dict) -> None:
    profile = org_profile or {}
    system_name = profile.get("system_name") or "Information System"
    system_desc = profile.get("system_description") or ""
    architecture = profile.get("architecture_summary") or ""
    boundary = profile.get("boundary_description") or ""
    uid = profile.get("system_unique_id") or "[Insert the System Unique Identifier]"

    for para in doc.paragraphs:
        t = para.text
        if "System Name/Title:" in t and "[" in t:
            para.text = f"System Name/Title: {system_name}"
        elif "System Unique Identifier:" in t:
            para.text = f"System Unique Identifier: {uid if uid != '[Insert the System Unique Identifier]' else 'TBD'}"
        elif "General Description/Purpose of System:" in t and system_desc:
            para.text = (
                "General Description/Purpose of System:  What is the function/purpose of the system?  "
                f"{system_desc}"
            )
        elif t.startswith("[Insert a system topology"):
            para.text = architecture or boundary or t
        elif "Number of end users" in t:
            cui = asset_scope.get("CUI Assets", 0)
            para.text = (
                "Number of end users and privileged users: "
                f"Approximately {cui or 'TBD'}% of assets process CUI — update user counts in the table below."
            )


def _fill_environment_inventories(
    doc, org_profile: dict, org_inventory: dict | None = None
) -> None:
    profile = org_profile or {}
    hardware = (profile.get("hardware_inventory") or "").strip()
    software = (profile.get("software_inventory") or "").strip()
    owned = (profile.get("hw_sw_org_owned") or "Yes").strip()
    structured = inventory_summary(org_inventory)

    for para in doc.paragraphs:
        t = para.text
        if "listing of all hardware" in t:
            extra = structured or hardware
            if extra:
                para.text = f"{t.split('[')[0].strip()} {extra}"
        elif "List all software components" in t:
            extra = software or (structured if structured and not hardware else "")
            if extra:
                para.text = f"List all software components installed on the system.  {extra}"
        elif "Hardware and Software Maintenance and Ownership" in t:
            if owned.lower().startswith("n"):
                para.text = f"{t.split('[Yes/No')[0].strip()} No — {owned}"
            else:
                para.text = f"{t.split('[Yes/No')[0].strip()} Yes"


def _insert_topology_diagram(doc, image_bytes: bytes | None) -> None:
    if not image_bytes:
        return
    from docx.oxml import OxmlElement
    from docx.shared import Inches
    from docx.text.paragraph import Paragraph

    anchor = None
    for i, para in enumerate(doc.paragraphs):
        if "Include a detailed topology narrative" in para.text:
            for j in range(i + 1, min(i + 6, len(doc.paragraphs))):
                if doc.paragraphs[j].text.strip():
                    anchor = doc.paragraphs[j]
                    break
            break
    if anchor is None:
        return

    new_p = OxmlElement("w:p")
    anchor._element.addnext(new_p)
    img_para = Paragraph(new_p, anchor._parent)
    img_para.alignment = 1  # center
    run = img_para.add_run()
    run.add_picture(BytesIO(image_bytes), width=Inches(6.25))


def _append_supporting_documents(doc, org_assets: dict | None) -> None:
    files = (org_assets or {}).get("appendix_files") or []
    if not files:
        return
    doc.add_page_break()
    doc.add_paragraph("APPENDIX — SUPPORTING DOCUMENTS")
    doc.add_paragraph(
        "The following documents support this SSP. Files are stored with the assessment workspace "
        "and are available to assessors upon request. Embed PDFs in this Word file manually if required."
    )
    for item in files:
        title = item.get("label") or item.get("filename", "Document")
        fname = item.get("filename", "")
        doc.add_paragraph(f"{title} — file: {fname}", style="List Paragraph")


def _fill_org_tables(doc, org_profile: dict) -> None:
    profile = org_profile or {}
    org_name = profile.get("org_name") or ""
    org_address = profile.get("org_address") or ""
    org_phone = profile.get("org_phone") or ""

    owner_name, owner_title = _split_name_title(profile.get("system_owner", ""))
    iso_name, iso_title = _split_name_title(profile.get("iso_name", ""))

    if len(doc.tables) < 4:
        return

    _set_cell(doc.tables[0], 0, 1, org_name)
    _set_cell(doc.tables[0], 1, 1, org_address)
    _set_cell(doc.tables[0], 2, 1, org_phone)

    # Table 1 — Information Owner (often N/A for contractors; leave labels)
    # Table 2 — System Owner
    _set_cell(doc.tables[2], 0, 1, owner_name)
    _set_cell(doc.tables[2], 1, 1, owner_title)

    # Table 3 — System Security Officer
    _set_cell(doc.tables[3], 0, 1, iso_name)
    _set_cell(doc.tables[3], 1, 1, iso_title)


def _implementation_tables(doc) -> list:
    return [
        t
        for t in doc.tables
        if t.rows
        and t.rows[0].cells
        and "Implemented" in t.rows[0].cells[0].text
    ]


def _requirement_paragraphs(doc) -> list:
    started = False
    reqs = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if "REQUIREMENTS" in t and para.style.name == "List Paragraph":
            started = True
            continue
        if not started:
            continue
        if para.style.name not in ("Header", "List Paragraph") or len(t) < 25:
            continue
        if t in _FAMILY_HEADERS:
            continue
        if t.startswith("Include ") or t.startswith("List all ") or "Number of end users" in t:
            continue
        if "Note:" in t[:10]:
            continue
        reqs.append(para)
    return reqs


def _set_form_checkbox_checked(cell, checked: bool) -> None:
    """Toggle legacy Word FORMCHECKBOX in a table header cell."""
    checkboxes = cell._tc.xpath(".//w:checkBox")
    if not checkboxes:
        return
    checkbox = checkboxes[0]
    for el in list(checkbox):
        if el.tag == qn("w:checked"):
            checkbox.remove(el)
    if checked:
        checked_el = OxmlElement("w:checked")
        checked_el.set(qn("w:val"), "1")
        checkbox.append(checked_el)


def _implementation_column_index(status: str) -> int:
    """
    Map assessment status to CUI template column:
    0 = Implemented, 1 = Planned to be Implemented, 2 = Not Applicable.
    """
    status = (status or "NOT STARTED").upper()
    if status == "MET":
        return 0
    if status in ("NOT APPLICABLE", "INHERITED"):
        return 2
    return 1


def _narrative_for_cell(status: str, narrative: str, cid: str, poam_id: str | None = None) -> tuple[int, str]:
    col_idx = _implementation_column_index(status)
    text = narrative.strip() or CMMC_FRAMEWORK.get(cid, {}).get("description", "")
    if col_idx == 0:
        return col_idx, text
    if col_idx == 2:
        return col_idx, text or "Not applicable to this system boundary."
    ref = f"POA&M Item {poam_id}" if poam_id else "the POA&M"
    if narrative.strip():
        return col_idx, f"{text} (Remediation tracked in {ref}.)"
    return col_idx, f"Planned — see {ref} for remediation details."


def _fill_implementation_tables(doc, answers: dict, scoped_controls: list) -> None:
    from poam_export import poam_weakness_ids

    impl_tables = _implementation_tables(doc)
    req_paras = _requirement_paragraphs(doc)
    scoped = set(scoped_controls)
    poam_ids = poam_weakness_ids(answers, scoped_controls)

    for idx, table in enumerate(impl_tables):
        cid = None
        if idx < len(req_paras):
            cid = _match_control_id(req_paras[idx].text)
        if not cid and idx < len(scoped_controls):
            # Last-resort: not reliable; prefer text match
            pass
        if not cid:
            continue

        ans = answers.get(cid, {})
        status = ans.get("status", "NOT STARTED")
        narrative = (
            ans.get("implementation_narrative", "")
            or ans.get("implementation_desc", "")
            or ans.get("assessor_notes", "")
        )

        if cid not in scoped:
            col_idx = 2
            text = "Out of scope for this assessment boundary."
        else:
            col_idx, text = _narrative_for_cell(status, narrative, cid, poam_ids.get(cid))

        supplement = control_ssp_supplement(ans)
        if supplement:
            text = f"{text}\n\n{supplement}".strip()

        if table.rows:
            header = table.rows[0]
            for ci in range(min(3, len(header.cells))):
                _set_form_checkbox_checked(header.cells[ci], ci == col_idx)

        # Row 1 cells are merged in the official template — write once.
        _set_cell(table, 1, 0, text[:8000])


def fill_cui_ssp_template(
    doc,
    answers: dict,
    asset_scope: dict,
    scoped_controls: list,
    org_profile: dict | None = None,
    org_assets: dict | None = None,
    org_asset_bytes: dict | None = None,
    audit_log: list | None = None,
    org_inventory: dict | None = None,
    risks_by_control: dict | None = None,
    system_scope: dict | None = None,
    ssp_version: str = "",
    prepared_by: str = "",
    approved_by: str = "",
) -> None:
    profile = org_profile or {}
    org_name = profile.get("org_name") or "Organization"
    fix_page_headers(doc, org_name, org_profile=profile)
    _fill_identification_paragraphs(doc, profile, asset_scope or {})
    _fill_org_tables(doc, profile)
    _fill_environment_inventories(doc, profile, org_inventory)
    topo = get_topology_bytes(org_assets or {}, org_asset_bytes or {})
    if topo:
        _insert_topology_diagram(doc, topo[1])
    _fill_implementation_tables(doc, answers, scoped_controls)
    _append_supporting_documents(doc, org_assets)
    add_asset_inventory_appendix(doc, org_inventory)
    add_record_of_changes(doc, audit_log)
    append_workspace_snapshot_notice(doc)


def template_control_mapping(doc) -> List[str]:
    """Debug: control IDs in template table order."""
    impl = _implementation_tables(doc)
    reqs = _requirement_paragraphs(doc)
    out = []
    for idx in range(len(impl)):
        cid = _match_control_id(reqs[idx].text) if idx < len(reqs) else None
        out.append(cid or "?")
    return out
